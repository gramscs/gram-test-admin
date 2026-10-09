"""Spreadsheet, CSV and management PDF exports of the same filtered snapshot."""
import csv
import io
import math
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.graphics.shapes import Drawing, Rect, String, Line, Circle
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

from app.admin.reporting import COLUMNS

MIMES = {'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'csv': 'text/csv; charset=utf-8', 'pdf': 'application/pdf'}
PALETTE = {'Pickup Scheduled': '#d3a14b', 'In Transit': '#6696ef', 'Out for Delivery': '#9774df', 'Delivered': '#358e71', 'No status': '#9ba7b0', 'Other': '#de7980'}


def text_cell(sheet, values):
    sheet.append(values)
    # Force user-entered identifiers, addresses and notes to remain text, never formulas.
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = 's'


def metric_rows(report):
    m = report['metrics']
    return [('Shipments', m['total']), ('Delivered (current status)', m['delivered']), ('Open shipments', m['open']), ('Delivered share (%)', m['delivered_share']), ('Pieces', m['pieces']), ('Chargeable weight (kg)', m['weight']), ('Chargeable volume (m³)', m['volume']), ('Weight recorded (shipments)', m['weight_recorded']), ('Volume recorded (shipments)', m['volume_recorded']), ('Open past expected date', m['past_expected']), ('Delivered without POD reference', m['missing_pod']), ('Without invoice reference', m['missing_invoice']), ('Unassigned client', m['unassigned']), ('Pickup date missing/invalid (included)', m['undated']), ('Undated excluded by date filter', m['excluded_undated']), ('Invalid measurements excluded from totals (shipments)', m['invalid_measurements'])]


def excel(report, columns):
    book = Workbook()
    summary = book.active
    summary.title = 'Management summary'
    for row in [('GRAM SCS · Shipment MIS', ''), ('Generated', report['generated_at']), ('Pickup-date period', report['period_label']), ('Client', report['client_label']), ('Status', report['status_label']), ('Definition', report['definition']), ('Measure', 'Value'), *metric_rows(report)]:
        text_cell(summary, row)
    sheets = [
        ('Status', ['Current status', 'Shipments'], [[x['label'], x['count']] for x in report['statuses']]),
        ('Pickup trend', [f'Pickup period ({report["grain"]})', 'Shipments'], [[x['label'], x['count']] for x in report['trend']]),
        ('Clients', ['Client', 'Shipments', 'Delivered', 'Pieces', 'Weight (kg)'], [[x['label'], x['count'], x['delivered'], x['pieces'], x['weight']] for x in report['clients']]),
        ('Routes', ['Pickup → drop', 'Shipments'], [[x['label'], x['count']] for x in report['routes']]),
        ('Attention', ['Shipment identifier', 'Client', 'Attention flags'], [[x['identifier'], x['client'], '; '.join(x['issues'])] for x in report['exceptions']]),
        ('Shipments', [COLUMNS[key] for key in columns], [[row[key] for key in columns] for row in report['details']]),
    ]
    for title, headers, rows in sheets:
        sheet = book.create_sheet(title)
        text_cell(sheet, headers)
        for row in rows:
            text_cell(sheet, row)
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='183B33')
            cell.font = Font(color='FFFFFF', bold=True)
        for index, header in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(55, max(20, len(header)+3))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
    summary.column_dimensions['A'].width = 48
    summary.column_dimensions['B'].width = 65
    for row in summary:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    summary.row_dimensions[6].height = 60
    summary['A1'].font = Font(size=19, bold=True, color='183B33')
    summary.freeze_panes = 'A8'
    for name, chart_type, anchor, title in [('Status', DoughnutChart, 'D2', 'Current status'), ('Pickup trend', LineChart, 'D18', 'Pickup volume'), ('Clients', BarChart, 'D34', 'Shipments by client')]:
        source = book[name]
        if source.max_row > 1:
            chart = chart_type()
            chart.title = title
            chart.add_data(Reference(source, min_col=2, max_col=2, min_row=1, max_row=source.max_row), titles_from_data=True)
            chart.set_categories(Reference(source, min_col=1, min_row=2, max_row=source.max_row))
            chart.width, chart.height = 21, 8
            if isinstance(chart, BarChart):
                chart.type = 'bar'
                chart.series[0].graphicalProperties.solidFill = '358E71'
            elif isinstance(chart, DoughnutChart):
                points = []
                for index, item in enumerate(report['statuses']):
                    point = DataPoint(idx=index)
                    point.graphicalProperties.solidFill = PALETTE[item['label']][1:]
                    points.append(point)
                chart.series[0].data_points = points
            else:
                chart.series[0].graphicalProperties.line.solidFill = '358E71'
            summary.add_chart(chart, anchor)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def csv_export(report, columns):
    buffer = io.StringIO(newline='')
    writer = csv.writer(buffer)
    writer.writerow([COLUMNS[key] for key in columns])
    for row in report['details']:
        values = []
        for key in columns:
            value = row[key]
            if isinstance(value, str) and (value.lstrip().startswith(('=', '+', '-', '@')) or value.startswith(('\t', '\r', '\n'))):
                value = "'" + value
            values.append(value)
        writer.writerow(values)
    return ('\ufeff' + buffer.getvalue()).encode('utf-8')


def pdf(report, columns):
    # Import registers the bundled embedded fonts, including on headless Linux.
    from app.admin import labels  # noqa: F401
    buffer = io.BytesIO()
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = 'LabelSans'
    styles.add(ParagraphStyle('MISHeading', fontName='LabelSansBold', fontSize=24, leading=30, textColor=colors.HexColor('#183B33'), spaceAfter=8))
    styles.add(ParagraphStyle('MISSmall', fontName='LabelSans', fontSize=9, leading=13, textColor=colors.HexColor('#536961')))
    def p(value, style='MISSmall'):
        return Paragraph(escape(str(value)), styles[style])
    story = [p('Shipment performance · Management MIS', 'MISHeading'), p(f'{report["period_label"]} · {report["client_label"]} · {report["status_label"]}'), p('Snapshot generated ' + report['generated_at']), Spacer(1, 16)]
    m = report['metrics']
    kpis = Table([[p('SHIPMENTS'), p('DELIVERED SHARE'), p('PIECES'), p('CHARGEABLE WEIGHT')], [p(f'{m["total"]:,}', 'MISHeading'), p(f'{m["delivered_share"]}%', 'MISHeading'), p(f'{m["pieces"]:,}', 'MISHeading'), p(f'{m["weight"]:,.3f} kg', 'MISHeading')]], colWidths=[189]*4)
    kpis.setStyle(TableStyle([('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#EDF5F1')), ('TOPPADDING',(0,0),(-1,-1),12), ('BOTTOMPADDING',(0,0),(-1,-1),8)]))
    story += [kpis, Spacer(1, 20)]
    drawing = Drawing(756, 180)
    drawing.add(String(0,165,'CURRENT DELIVERY STATUS',fontName='LabelSansBold',fontSize=11,fillColor=colors.HexColor('#183B33')))
    y = 140
    for item in report['statuses']:
        drawing.add(String(0,y,item['label'],fontName='LabelSans',fontSize=9))
        drawing.add(Rect(118,y-2,max(1,item['count']/max(m['total'],1)*160),9,fillColor=colors.HexColor(PALETTE[item['label']]),strokeColor=None))
        drawing.add(String(290,y,str(item['count']),fontName='LabelSansBold',fontSize=10))
        y -= 23
    drawing.add(String(350,165,'PICKUP VOLUME',fontName='LabelSansBold',fontSize=11,fillColor=colors.HexColor('#183B33')))
    trend = report['trend']
    if trend:
        peak = max(x['count'] for x in trend)
        high = max(1, math.ceil(peak / 4)) * 4
        for tick in range(5):
            axis_y = 32 + tick * 110 / 4
            drawing.add(Line(370,axis_y,735,axis_y,strokeColor=colors.HexColor('#DFE8E3'),strokeWidth=.5))
            drawing.add(String(350,axis_y-3,str(high*tick//4),fontName='LabelSans',fontSize=8,fillColor=colors.HexColor('#536961')))
        prev = None
        for index, item in enumerate(trend):
            x, y = 370 + index/max(len(trend)-1,1)*365, 32+item['count']/high*110
            if prev:
                drawing.add(Line(*prev,x,y,strokeColor=colors.HexColor('#23886c'),strokeWidth=2))
            if len(trend)<31:
                drawing.add(Circle(x,y,2,fillColor=colors.HexColor('#23886c'),strokeColor=None))
            prev = (x,y)
        drawing.add(String(350,12,trend[0]['label'],fontName='LabelSans',fontSize=8))
        drawing.add(String(665,12,trend[-1]['label'],fontName='LabelSans',fontSize=8))
        drawing.add(String(350,148,f'Peak: {peak} shipments per {report["grain"][:-2] if report["grain"] != "daily" else "day"}',fontName='LabelSans',fontSize=8))
    else:
        drawing.add(String(350,95,'No valid pickup dates in this view',fontName='LabelSans',fontSize=10))
    watchlist = Table([[p('OPEN PAST EXPECTED DATE'), p('DELIVERED WITHOUT POD'), p('INVOICES MISSING')], [p(m['past_expected'], 'Heading2'), p(m['missing_pod'], 'Heading2'), p(m['missing_invoice'], 'Heading2')]], colWidths=[252]*3)
    watchlist.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#F6F8F7')),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
    story += [drawing, watchlist, Spacer(1,10), p(report['definition']), PageBreak()]
    def table(title, headers, rows, widths):
        story.append(p(title, 'Heading2'))
        content = [[p(value) for value in headers]] + [[p(value) for value in row] for row in rows]
        result = Table(content, colWidths=widths, repeatRows=1, hAlign='LEFT')
        result.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E7F0EC')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F6F8F7')]),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
        story.extend([result, Spacer(1,12)])
    table('Operational attention & data coverage', ['Measure', 'Value'], metric_rows(report)[9:], [590,166])
    table('Client performance · top 15 by shipment volume', ['Client','Shipments','Delivered','Pieces','Weight (kg)'], [[x['label'],x['count'],x['delivered'],x['pieces'],f'{x["weight"]:,.3f}'] for x in report['clients'][:15]] or [['No shipments','0','0','0','0']], [350,90,90,90,136])
    table('Routes · top 15 by shipment volume', ['Pickup → drop','Shipments'], [[x['label'],x['count']] for x in report['routes'][:15]] or [['No routes',0]], [650,106])
    table(f'Attention register · first 50 of {len(report["exceptions"])} shipments', ['Shipment','Client','Attention flags'], [[x['identifier'],x['client'],'; '.join(x['issues'])] for x in report['exceptions'][:50]] or [['No attention flags','','']], [150,200,406])
    story.append(p('Excel and CSV contain the full filtered shipment register. PDF client/route tables show the top 15; the attention register shows the first 50 in shipment record order. Blank measurements are excluded from totals.'))
    def footer(canvas, document):
        canvas.setFont('LabelSans',8)
        canvas.setFillColor(colors.HexColor('#536961'))
        canvas.drawString(42,22,'GRAM SCS · Shipment MIS · '+report['generated_at'])
        canvas.drawRightString(800,22,str(document.page))
    SimpleDocTemplate(buffer,pagesize=landscape(A4),leftMargin=42,rightMargin=42,topMargin=34,bottomMargin=40,title='GRAM SCS Shipment MIS',author='GRAM SCS').build(story,onFirstPage=footer,onLaterPages=footer)
    return buffer.getvalue()


def export_bytes(report, columns, output_format):
    if output_format not in MIMES:
        raise ValueError('Choose Excel, PDF or CSV.')
    return {'xlsx': excel, 'pdf': pdf, 'csv': csv_export}[output_format](report, columns)
