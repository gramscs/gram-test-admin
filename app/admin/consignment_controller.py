"""Admin consignment management routes and helpers."""

from decimal import Decimal, InvalidOperation
import io
import logging
import os
import re
from pathlib import Path

from flask import current_app, flash, jsonify, redirect, render_template, request, send_file, url_for
from openpyxl import Workbook, load_workbook
from sqlalchemy.exc import DatabaseError, OperationalError, ProgrammingError

from app import limiter
from app.admin import admin_bp
from app.admin.auth import require_admin
from app.models import db
from app import orm as models
from app.admin.documents import prepare_row_documents, store_document_changes, cleanup_documents

logger = logging.getLogger(__name__)



def _get_supabase_client():
    url = os.getenv("SUPABASE_URL", "").strip()
    key = os.getenv("SUPABASE_KEY", "").strip()
    if not url or not key:
        return None
    try:
        from supabase import create_client
    except Exception:
        logger.warning("Supabase package not available; falling back to local uploads")
        return None
    try:
        return create_client(url, key)
    except Exception:
        logger.exception("Failed to create Supabase client")
        return None


def _store_pod_bytes(filename, file_bytes, content_type=None, bucket_name=None):
    supa = _get_supabase_client()
    if supa:
        bucket = bucket_name or os.getenv("SUPABASE_BUCKET", "pod-uploads")
        object_path = f"consignments/{filename}"
        payload = file_bytes
        if hasattr(payload, "read"):
            payload = payload.read()
        if isinstance(payload, bytearray):
            payload = bytes(payload)
        if not isinstance(payload, (bytes, bytearray)):
            raise TypeError("POD upload payload must be bytes-like.")
        try:
            supa.storage.from_(bucket).upload(object_path, payload, {"content-type": content_type or "application/octet-stream"})
        except Exception:
            try:
                supa.storage.from_(bucket).remove([object_path])
            except Exception:
                logger.exception("Failed to clean incomplete document upload")
            raise
        return f"supabase:{bucket}/{object_path}"

    payload = file_bytes
    if hasattr(payload, "read"):
        payload = payload.read()
    if isinstance(payload, bytearray):
        payload = bytes(payload)
    if not isinstance(payload, (bytes, bytearray)):
        raise TypeError("POD upload payload must be bytes-like.")

    upload_folder = os.path.join(current_app.instance_path, "uploads")
    os.makedirs(upload_folder, exist_ok=True)
    dest_path = os.path.join(upload_folder, filename)
    opened = False
    try:
        with open(dest_path, "xb") as handle:
            opened = True
            handle.write(payload)
    except Exception:
        if opened and os.path.isfile(dest_path):
            os.remove(dest_path)
        raise
    return filename


def _parse_supabase_pod_value(pod_value):
    if not isinstance(pod_value, str) or not pod_value.startswith('supabase:'):
        raise ValueError('POD is not stored in Supabase.')

    _, rest = pod_value.split(':', 1)
    bucket, object_path = rest.split('/', 1)
    if not bucket or not object_path:
        raise ValueError('Invalid Supabase POD path.')

    return bucket, object_path


def _download_supabase_pod_file(pod_value):
    client = _get_supabase_client()
    if not client:
        raise RuntimeError('Supabase not configured.')

    bucket, object_path = _parse_supabase_pod_value(pod_value)
    content = client.storage.from_(bucket).download(object_path)
    if hasattr(content, 'read'):
        content = content.read()
    if isinstance(content, bytearray):
        content = bytes(content)
    if not isinstance(content, bytes):
        raise RuntimeError('Unexpected Supabase download response.')

    return content, object_path


def _download_legacy_supabase_pod_file(consignment_id, pod_value):
    """Download a legacy POD by attempting old Supabase object paths.

    Legacy records may store a bare object path or a local filename that was
    previously migrated into Supabase. This helper tries the configured bucket
    and an optional consignment-id-prefixed path.
    """
    client = _get_supabase_client()
    if not client:
        raise RuntimeError('Supabase not configured.')

    bucket = os.getenv('SUPABASE_BUCKET', 'pod-uploads')
    if not isinstance(pod_value, str) or not pod_value:
        raise ValueError('Invalid legacy POD path.')

    candidates = []
    legacy_bucket = bucket
    if pod_value.startswith('supabase:'):
        _, rest = pod_value.split(':', 1)
        try:
            legacy_bucket, legacy_object_path = rest.split('/', 1)
        except ValueError:
            legacy_bucket = bucket
            legacy_object_path = rest
        candidates.append((legacy_bucket, legacy_object_path))
    else:
        candidates.append((bucket, pod_value))
        if consignment_id is not None:
            consignment_prefix = str(consignment_id)
            if not pod_value.startswith(consignment_prefix + '/'):
                candidates.append((bucket, f"{consignment_prefix}/{pod_value}"))

    last_error = None
    for candidate_bucket, object_path in candidates:
        try:
            content = client.storage.from_(candidate_bucket).download(object_path)
            if hasattr(content, 'read'):
                content = content.read()
            if isinstance(content, bytearray):
                content = bytes(content)
            if not isinstance(content, bytes):
                raise RuntimeError('Unexpected Supabase download response.')
            return content, candidate_bucket, object_path
        except Exception as exc:
            last_error = exc

    # Last chance: if the value was a local filename under uploads, try that path.
    upload_folder = os.path.join(current_app.instance_path, 'uploads')
    try:
        legacy_path = os.path.normpath(os.path.join(upload_folder, pod_value))
        if Path(legacy_path).resolve().is_relative_to(Path(upload_folder).resolve()) and os.path.isfile(legacy_path):
            with open(legacy_path, 'rb') as fh:
                return fh.read(), bucket, pod_value
    except Exception:
        pass

    raise RuntimeError('Legacy POD file not found.') from last_error


def _delete_pod_file(pod_value):
    if not pod_value:
        return

    if isinstance(pod_value, str) and pod_value.startswith("supabase:"):
        client = _get_supabase_client()
        if not client:
            return
        try:
            _, rest = pod_value.split(":", 1)
            bucket, object_path = rest.split("/", 1)
            client.storage.from_(bucket).remove([object_path])
        except Exception:
            logger.exception("Failed to remove POD from Supabase")
        return

    upload_folder = os.path.join(current_app.instance_path, "uploads")
    pod_value = pod_value.removeprefix('local:')
    pod_path = os.path.normpath(os.path.join(upload_folder, pod_value))
    if Path(pod_path).resolve().is_relative_to(Path(upload_folder).resolve()) and os.path.isfile(pod_path):
        try:
            os.remove(pod_path)
        except Exception:
            logger.exception("Failed to remove POD file from disk")


def _is_external_pod_url(value):
    if not isinstance(value, str):
        return False
    lowered = value.strip().lower()
    return lowered.startswith("http://") or lowered.startswith("https://")


def _normalize_header(value):
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _serialize_consignment(consignment):
    return {
        "id": getattr(consignment, "id", None),
        "consignment_number": getattr(consignment, "consignment_number", None),
        "identifier_type": consignment.identifier_type or "LRN",
        "pieces": consignment.pieces or 1,
        "chargeable_weight": str(consignment.chargeable_weight) if consignment.chargeable_weight is not None else None,
        "chargeable_volume": str(consignment.chargeable_volume) if consignment.chargeable_volume is not None else None,
        "company_id": consignment.company_id,
        "company_name": consignment.company.name if consignment.company else None,
        "status": getattr(consignment, "status", None),
        "pickup_pincode": getattr(consignment, "pickup_pincode", None),
        "pickup_address": getattr(consignment, "pickup_address", None),
        "pickup_tag": getattr(consignment, "pickup_tag", None),
        "pickup_date": getattr(consignment, "pickup_date", None),
        "drop_pincode": getattr(consignment, "drop_pincode", None),
        "drop_address": getattr(consignment, "drop_address", None),
        "drop_tag": getattr(consignment, "drop_tag", None),
        "drop_date": getattr(consignment, "drop_date", None),
        "eta": getattr(consignment, "eta", None),
        "pod_image": getattr(consignment, "pod_image", None),
        "pod_original_name": consignment.pod_original_name,
        "invoice_file": consignment.invoice_file,
        "invoice_original_name": consignment.invoice_original_name,
        "pod_file_name": getattr(consignment, "pod_file_name", None),
        "pod_file_type": getattr(consignment, "pod_file_type", None),
        "pod_file_data": getattr(consignment, "pod_file_data", None),
    }


@admin_bp.route("/admin/consignments", methods=["GET"], endpoint="consignments_panel")
@require_admin
def consignments_panel():
    try:
        total = models.shipment_query().count()
        consignments = [] if total > 500 else models.shipment_query().order_by(models.Consignment.id.asc()).limit(200).all()
        rows = [_serialize_consignment(row) for row in consignments]
        return render_template("admin/consignments.html", consignments=rows)
    except (OperationalError, DatabaseError, ProgrammingError):
        logger.exception("Database error loading admin consignments panel")
        return render_template(
            "admin/consignments.html",
            consignments=[],
            error="Unable to load data. Please try again.",
        )
    except Exception:
        logger.exception("Unexpected error in admin consignments panel")
        return render_template(
            "admin/consignments.html",
            consignments=[],
            error="An unexpected error occurred.",
        )


@admin_bp.route("/admin/consignments/list", methods=["GET"], endpoint="consignments_list_api")
@require_admin
def consignments_list_api():
    try:
        page = max(1, request.args.get("page", 1, type=int))
        per_page = max(1, min(100, request.args.get("per_page", 10, type=int)))
        search = request.args.get("search", "", type=str).strip()
        sort_by = request.args.get("sort_by", "id", type=str)
        sort_order = request.args.get("sort_order", "asc", type=str)

        allowed_sort_columns = {
            "id", "consignment_number", "identifier_type", "pieces", "chargeable_weight", "chargeable_volume", "status", "pickup_pincode", "drop_pincode",
            "pickup_tag", "drop_tag", "pickup_date", "drop_date",
        }
        if sort_by not in allowed_sort_columns:
            sort_by = "id"
        sort_order = "asc" if sort_order.lower() == "asc" else "desc"

        from app.admin.reporting import search_shipments
        query = search_shipments(models.shipment_query(), search)

        total = query.count()
        sort_column = getattr(models.Consignment, sort_by)
        query = query.order_by(sort_column.desc() if sort_order == "desc" else sort_column.asc())
        rows = query.offset((page - 1) * per_page).limit(per_page).all()

        pages = (total + per_page - 1) // per_page if total else 0
        return jsonify({
            "success": True,
            "rows": [_serialize_consignment(row) for row in rows],
            "page": page,
            "per_page": per_page,
            "total": total,
            "pages": pages,
            "has_prev": page > 1,
            "has_next": page < pages,
        })
    except Exception as exc:
        logger.exception("Consignment list API failed: %s", exc)
        return jsonify({
            "success": False,
            "rows": [],
            "page": 1,
            "per_page": 10,
            "total": 0,
            "pages": 0,
            "has_prev": False,
            "has_next": False,
            "error": "Unable to load consignments right now.",
        }), 500


@admin_bp.route("/admin/consignments/import-template.xlsx", methods=["GET"], endpoint="consignments_import_template_excel")
@require_admin
def consignments_import_template_excel():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Consignments"
    sheet.append([
        "consignment_number",
        "status",
        "pickup_address",
        "pickup_pincode",
        "pickup_tag",
        "pickup_date",
        "drop_address",
        "drop_pincode",
        "drop_tag",
        "drop_date",
        "identifier_type", "pieces", "chargeable_weight", "chargeable_volume", "company_id",
    ])

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        as_attachment=True,
        download_name="consignment_import_template.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@admin_bp.route("/admin/consignments/import", methods=["POST"], endpoint="consignments_import_excel")
@limiter.limit("10 per minute")
@require_admin
def consignments_import_excel():
    uploaded_file = request.files.get("file")
    if not uploaded_file:
        flash("Please choose an Excel file to import.", "danger")
        return redirect(url_for("admin.consignments_panel"))

    workbook = load_workbook(uploaded_file, data_only=True)
    sheet = workbook.active
    headers = [_normalize_header(cell) for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]

    existing_numbers = {
        row[0]
        for row in models.shipment_query().with_entities(models.Consignment.consignment_number).all()
        if row and row[0]
    }

    added_count = 0
    skipped_duplicates = 0

    for row_values in sheet.iter_rows(min_row=2, values_only=True):
        row_data = {headers[index]: value for index, value in enumerate(row_values) if index < len(headers)}
        consignment_number = str(row_data.get("consignment_number") or "").strip()
        if not consignment_number:
            continue
        if consignment_number in existing_numbers:
            skipped_duplicates += 1
            continue

        consignment = models.Consignment(
            consignment_number=consignment_number,
            status=row_data.get("status") or "",
            pickup_address=row_data.get("pickup_address"),
            pickup_pincode=row_data.get("pickup_pincode"),
            pickup_tag=row_data.get("pickup_tag"),
            pickup_date=row_data.get("pickup_date"),
            drop_address=row_data.get("drop_address"),
            drop_pincode=row_data.get("drop_pincode"),
            drop_tag=row_data.get("drop_tag"),
            drop_date=row_data.get("drop_date"),
        )

        try:
            fields = _normalize_shipment_fields(row_data)
            for field, value in fields.items():
                setattr(consignment, field, value)
        except ValueError as exc:
            db.session.rollback()
            flash(f"Import failed for {consignment_number}: {exc}", "danger")
            return redirect(url_for("admin.consignments_panel"))
        db.session.add(consignment)
        existing_numbers.add(consignment_number)
        added_count += 1

    try:
        db.session.commit()
        flash(f"Import completed. Added: {added_count}, skipped duplicates: {skipped_duplicates}.", "success")
    except Exception:
        db.session.rollback()
        logger.exception("Failed to import consignments")
        flash("Import failed.", "danger")

    return redirect(url_for("admin.consignments_panel"))


@admin_bp.route("/admin/consignments/export.xlsx", methods=["GET"], endpoint="consignments_export_excel")
@require_admin
def consignments_export_excel():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Consignments"
    headers = [
        "consignment_number",
        "status",
        "pickup_tag",
        "drop_pincode",
        "pickup_date",
        "drop_date",
        "pickup_address",
        "drop_address",
        "identifier_type", "pieces", "chargeable_weight", "chargeable_volume", "company_id", "client_company",
    ]
    sheet.append(headers)

    rows = models.shipment_query().order_by(models.Consignment.id.asc()).all()
    for consignment in rows:
        sheet.append([
            getattr(consignment, "consignment_number", None),
            getattr(consignment, "status", None),
            getattr(consignment, "pickup_tag", None),
            getattr(consignment, "drop_pincode", None),
            getattr(consignment, "pickup_date", None),
            getattr(consignment, "drop_date", None),
            getattr(consignment, "pickup_address", None),
            getattr(consignment, "drop_address", None),
            consignment.identifier_type, consignment.pieces,
            consignment.chargeable_weight, consignment.chargeable_volume,
            models.export_id(consignment.company_id), consignment.company.name if consignment.company else None,
        ])

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        as_attachment=True,
        download_name="consignments.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@admin_bp.route("/admin/consignments/export.pdf", methods=["GET"], endpoint="consignments_export_pdf")
@limiter.limit("6 per minute")
@require_admin
def consignments_export_pdf():
    from app.admin.mis import export_response
    return export_response('pdf')


CONSIGNMENT_SAVE_FIELDS = (
    "consignment_number",
    "status",
    "pickup_pincode",
    "pickup_address",
    "pickup_tag",
    "pickup_date",
    "drop_pincode",
    "drop_address",
    "drop_tag",
    "drop_date",
    "eta",
)


def _normalize_shipment_fields(row, existing=None):
    if "consignment_number" in row:
        number = str(row["consignment_number"] or "").strip()
        if not number or len(number) > 64 or not number.isascii() or not all(32 <= ord(char) <= 126 for char in number):
            raise ValueError("Identifier must contain 1–64 printable ASCII characters.")
    def value(name, default):
        return row[name] if name in row else getattr(existing, name, default)

    identifier_type = str(value("identifier_type", "LRN") or "LRN").strip()
    if identifier_type not in {"LRN", "Order ID", "AWB"}:
        raise ValueError("Choose LRN, Order ID, or AWB.")
    raw_pieces = value("pieces", 1)
    try:
        pieces = Decimal(str(1 if raw_pieces in (None, "") else raw_pieces))
        if not pieces.is_finite() or pieces != pieces.to_integral_value() or not 1 <= pieces <= 10000:
            raise ValueError()
    except (InvalidOperation, ValueError):
        raise ValueError("Pieces must be a whole number from 1 to 10000.") from None
    fields = {"identifier_type": identifier_type, "pieces": int(pieces)}
    company_id = value("company_id", None)
    if company_id in (None, ""):
        fields["company_id"] = None
    else:
        from app.admin.input_validation import positive_id
        company_id = positive_id(company_id)
        company = db.session.get(models.Company, company_id)
        if not company or (not company.active and company_id != getattr(existing, "company_id", None)):
            raise ValueError("Choose an active client company. Existing archived links can be retained.")
        fields["company_id"] = company_id
    if models.using_er() and fields['company_id'] != getattr(existing, 'company_id', None):
        fields['pickup_location_id'] = None
        fields['drop_location_id'] = None
    for name in ("chargeable_weight", "chargeable_volume"):
        raw = value(name, None)
        if raw in (None, ""):
            fields[name] = None
            continue
        try:
            number = Decimal(str(raw))
            if not number.is_finite() or number < 0 or number > Decimal("999999999.999") or number != number.quantize(Decimal("0.001")):
                raise ValueError()
        except (InvalidOperation, ValueError):
            raise ValueError(f"{name.replace('_', ' ').title()} must be non-negative with up to 3 decimal places.") from None
        fields[name] = number
    return fields


def _apply_consignment_payload(consignment, row):
    for field in CONSIGNMENT_SAVE_FIELDS:
        value = row.get(field)
        if value is None:
            value = ""
        if isinstance(value, str):
            value = value.strip()
        setattr(consignment, field, value)


def _normalize_save_payload():
    if request.is_json:
        payload = request.get_json(silent=True) or {}
    else:
        payload = request.form.to_dict(flat=True)

    if isinstance(payload.get("rows"), list):
        return payload

    # Backwards-compatible single-row payload support.
    if payload.get("consignment_number") or payload.get("id"):
        return {"rows": [payload], "deleted_ids": []}

    return payload


@admin_bp.route("/admin/consignments/save", methods=["POST"], endpoint="consignments_save")
@require_admin
def consignments_save():
    payload = _normalize_save_payload()
    rows = payload.get("rows") if isinstance(payload, dict) else None
    deleted_ids = payload.get("deleted_ids", []) if isinstance(payload, dict) else []

    if not isinstance(rows, list):
        return jsonify({"success": False, "message": "Rows payload is required."}), 400

    errors = []
    saved_count = 0
    deleted_count = 0
    document_plans, created_files, obsolete_files = [], [], []

    try:
        for deleted_id in deleted_ids if isinstance(deleted_ids, list) else []:
            try:
                deleted_id_int = models.record_id(deleted_id)
            except (TypeError, ValueError):
                continue
            consignment = db.session.get(models.Consignment, deleted_id_int)
            if consignment:
                obsolete_files.extend(value for value in (consignment.pod_image, consignment.invoice_file) if value)
                db.session.delete(consignment)
                deleted_count += 1

        for index, row in enumerate(rows):
            if not isinstance(row, dict):
                errors.append({"index": index, "field": "row", "message": "Row must be an object."})
                continue

            consignment_number = (row.get("consignment_number") or "").strip().upper()
            if not consignment_number:
                errors.append({
                    "index": index,
                    "field": "consignment_number",
                    "message": "Consignment number is required.",
                })
                continue

            row = dict(row)
            row["consignment_number"] = consignment_number

            consignment = None
            row_id = row.get("id")
            # Older browser sessions send negative IDs for unsaved staged rows.
            if str(row_id).startswith('-') and str(row_id)[1:].isdigit():
                row_id = None
            if row_id not in (None, ""):
                try:
                    consignment = db.session.get(models.Consignment, models.record_id(row_id))
                except (TypeError, ValueError):
                    errors.append({"index": index, "field": "id", "message": "Invalid consignment id."})
                    continue

            existing_by_number = models.shipment_query().filter_by(consignment_number=consignment_number).first()
            if consignment and existing_by_number and existing_by_number.id != consignment.id:
                errors.append({
                    "index": index,
                    "field": "consignment_number",
                    "message": "Consignment number already exists.",
                })
                continue

            if not consignment:
                consignment = existing_by_number or models.Consignment()

            if len(consignment_number) > 64 or not consignment_number.isascii() or not all(32 <= ord(char) <= 126 for char in consignment_number):
                errors.append({"index": index, "field": "consignment_number", "message": "Identifier must contain 1–64 printable ASCII characters."})
                continue
            try:
                fields = _normalize_shipment_fields(row, consignment)
                document_changes = prepare_row_documents(row, consignment)
            except ValueError as exc:
                errors.append({"index": index, "field": "shipment", "message": str(exc)})
                continue
            _apply_consignment_payload(consignment, row)
            for field, value in fields.items():
                setattr(consignment, field, value)
            db.session.add(consignment)
            document_plans.append((consignment, document_changes))
            saved_count += 1

        if errors:
            db.session.rollback()
            return jsonify({
                "success": False,
                "message": "Validation errors. Please fix highlighted rows.",
                "errors": errors,
            }), 400

        for consignment, changes in document_plans:
            store_document_changes(consignment, changes, created_files, obsolete_files)
        total = models.shipment_query().count()
        db.session.commit()
        cleanup_documents(obsolete_files)
        return jsonify({
            "success": True,
            "message": "Saved.",
            "saved_count": saved_count,
            "deleted_count": deleted_count,
            "total": total,
        })
    except Exception:
        db.session.rollback()
        cleanup_documents(created_files)
        logger.exception("Failed to save consignments")
        return jsonify({"success": False, "message": "Failed to save consignments."}), 500
