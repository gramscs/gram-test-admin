import io
from email import policy
from email.parser import BytesParser

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import create_engine, text

from app import create_app
from app.models import Company, Consignment, Lead, db


def company(client, name='Acme Logistics', **changes):
    data = {'name': name, 'address': 'Registered office', 'locations': [
        {'kind': 'pickup', 'label': 'Delhi warehouse', 'address': 'Sender Road', 'pincode': '110001', 'is_default': True},
        {'kind': 'drop', 'label': 'Mumbai store', 'address': 'Receiver Road', 'pincode': '400001', 'is_default': True}], **changes}
    response = client.post('/admin/companies', json=data)
    assert response.status_code == 200, response.get_json()
    return response.get_json()['company']


def enquiry(client, name='Test customer', **changes):
    response = client.post('/admin/leads', json={'name': name, 'email': 'customer@example.test', 'phone': '+91 9000000000', 'subject': 'Quotation', 'message': 'Please quote shipping.', **changes})
    assert response.status_code == 200, response.get_json()
    return response.get_json()['lead']


def test_company_presets_defaults_edit_and_shipment_snapshots(admin_client):
    master = company(admin_client)
    assert len(master['locations']) == 2
    assert all(row['is_default'] for row in master['locations'])
    shipment = {'consignment_number': 'CLIENT001', 'company_id': master['id'], 'pickup_address': 'Custom sender', 'pickup_pincode': '110001', 'drop_address': 'Receiver Road', 'drop_pincode': '400001'}
    assert admin_client.post('/admin/consignments/save', json={'rows': [shipment]}).status_code == 200
    master['locations'][0]['address'] = 'Changed master address'
    assert admin_client.put(f"/admin/companies/{master['id']}", json=master).status_code == 200
    saved = admin_client.get('/admin/consignments/list?search=Acme').get_json()['rows'][0]
    assert saved['company_id'] == master['id'] and saved['company_name'] == master['name']
    assert saved['pickup_address'] == 'Custom sender'
    assert admin_client.post(f"/admin/companies/{master['id']}/archive", json={}).status_code == 200
    assert admin_client.post('/admin/consignments/save', json={'rows': [saved]}).status_code == 200
    assert admin_client.post('/admin/consignments/save', json={'rows': [{'consignment_number': 'ARCHIVEDNEW', 'company_id': master['id']}]}).status_code == 400
    master['active'] = True
    assert admin_client.put(f"/admin/companies/{master['id']}", json=master).status_code == 200
    assert admin_client.post('/admin/consignments/save', json={'rows': [{'consignment_number': 'BADCLIENT', 'company_id': 999}]}).status_code == 400
    assert admin_client.get('/admin/consignments/list').get_json()['total'] == 1
    backup = admin_client.get('/admin/generate-backup').get_json()
    assert backup['companies'][0]['name'] == master['name']
    assert backup['company_locations'][0]['address'] == 'Changed master address'


def test_company_validation_and_foreign_location_rollback(admin_client):
    master = company(admin_client)
    assert admin_client.post('/admin/companies', json={'name': 'acme logistics'}).status_code == 400
    other = company(admin_client, name='Other company')
    master['locations'][0]['id'] = other['locations'][0]['id']
    master['name'] = 'Should not save'
    assert admin_client.put(f"/admin/companies/{master['id']}", json=master).status_code == 400
    options = admin_client.get('/admin/companies/options').get_json()['companies']
    assert options[0]['name'] == 'Acme Logistics'
    assert admin_client.post('/admin/companies', json={'name': 'Broken', 'locations': [{'kind': 'pickup', 'label': 'Bad PIN', 'address': 'Address', 'pincode': 'bad'}]}).status_code == 400
    assert len(admin_client.get('/admin/companies/options').get_json()['companies']) == 2


def test_client_link_excel_roundtrip_and_legacy_update(admin_client):
    master = company(admin_client)
    admin_client.post('/admin/consignments/save', json={'rows': [{'consignment_number': 'EXCLIENT', 'company_id': master['id']}]})
    exported = load_workbook(io.BytesIO(admin_client.get('/admin/consignments/export.xlsx').data)).active
    values = dict(zip([cell.value for cell in exported[1]], [cell.value for cell in exported[2]]))
    assert values['client_company'] == master['name'] and values['company_id'] == master['id']
    row = admin_client.get('/admin/consignments/list').get_json()['rows'][0]
    admin_client.post('/admin/consignments/save', json={'rows': [{'id': row['id'], 'consignment_number': row['consignment_number'], 'status': 'Delivered'}]})
    assert admin_client.get('/admin/consignments/list').get_json()['rows'][0]['company_id'] == master['id']


def test_enquiry_create_edit_filter_export_and_bulk_status(admin_client, app):
    row = enquiry(admin_client, notes='Internal only', company_name='Acme', follow_up_date='2026-10-15')
    old_time = row['created_at']
    result = admin_client.put(f"/admin/leads/{row['id']}", json={'phone': '9999999999', 'status': 'Qualified', 'message': '=FORMULA()', 'follow_up_date': '2026-10-20'}).get_json()['lead']
    assert result['name'] == row['name'] and result['created_at'] == old_time and result['notes'] == 'Internal only'
    assert result['follow_up_date'] == '2026-10-20'
    assert b'Test customer' in admin_client.get('/admin/leads?search=Acme&status=Qualified').data
    assert b'Test customer' not in admin_client.get('/admin/leads?status=New').data
    exported = load_workbook(io.BytesIO(admin_client.get(f"/admin/leads/export.xlsx?ids={row['id']}").data)).active
    message = exported.cell(2, 7)
    assert message.value == '=FORMULA()' and message.data_type == 's'
    other = enquiry(admin_client, name='Second customer')
    response = admin_client.post('/admin/leads/bulk', json={'ids': [row['id'], other['id']], 'action': 'status', 'status': 'Won'})
    assert response.get_json()['changed'] == 2
    assert b'Second customer' in admin_client.get('/admin/leads?status=Won').data
    assert admin_client.post('/admin/leads/bulk', json={'ids': [row['id'], 999], 'action': 'delete'}).status_code == 400
    with app.app_context():
        assert Lead.query.count() == 2
    assert admin_client.post('/admin/leads/bulk', json={'ids': [row['id'], other['id']], 'action': 'delete'}).get_json()['changed'] == 2


@pytest.mark.parametrize('changes', [{'email': 'invalid'}, {'name': ''}, {'status': 'Bad'}, {'follow_up_date': '2026-02-30'}, {'message': ''}, {'notes': 123}])
def test_bad_enquiry_does_not_persist(admin_client, app, changes):
    assert admin_client.post('/admin/leads', json={'name': 'Name', 'email': 'customer@example.test', 'message': 'Message', **changes}).status_code == 400
    with app.app_context():
        assert Lead.query.count() == 0


def test_email_preview_excludes_internal_notes_and_download_is_unsent(admin_client):
    row = enquiry(admin_client, message='Hello café', notes='Private team note', status='Qualified')
    preview = admin_client.post('/admin/leads/email-preview', json={'ids': [row['id']]}).get_json()
    assert 'Hello café' in preview['body'] and 'Private team note' not in preview['body'] and 'Status:' not in preview['body']
    with_notes = admin_client.post('/admin/leads/email-preview', json={'ids': [row['id']], 'include_notes': True}).get_json()
    assert 'Private team note' in with_notes['body'] and 'Qualified' in with_notes['body']
    response = admin_client.post('/admin/leads/email-draft', data={'ids': [row['id']], 'to': 'chosen@example.test,second@example.test', 'subject': 'Editable subject', 'body': preview['body']})
    assert response.status_code == 200 and response.mimetype == 'message/rfc822'
    draft = BytesParser(policy=policy.default).parsebytes(response.data)
    assert draft['To'] == 'chosen@example.test, second@example.test' and draft['Subject'] == 'Editable subject'
    assert draft['X-Unsent'] == '1' and 'Hello café' in draft.get_content()
    assert admin_client.post('/admin/leads/email-draft', data={'ids': [row['id']], 'to': 'invalid', 'subject': 'Subject', 'body': 'Body'}).status_code == 400
    assert admin_client.post('/admin/leads/email-draft', data={'ids': [row['id']], 'to': 'chosen@example.test', 'subject': 'Subject\r\nBcc: hidden@example.test', 'body': 'Body'}).status_code == 400
    assert admin_client.post('/admin/leads/email-preview', json={'ids': [999]}).status_code == 400


def test_updated_labels_have_optional_barcode_and_company_name(admin_client):
    master = company(admin_client)
    admin_client.post('/admin/consignments/save', json={'rows': [{'consignment_number': 'BARCODE001', 'company_id': master['id'], 'pieces': 2, 'pickup_address': 'Sender Road', 'drop_address': 'Receiver Road'}]})
    row = admin_client.get('/admin/consignments/list').get_json()['rows'][0]
    outputs = []
    for barcode in ['1', '0']:
        response = admin_client.post('/admin/labels/generate', data={'consignment_ids': [row['id']], 'barcode': barcode, 'delivery': 'preview'})
        assert response.status_code == 200 and response.headers['Content-Disposition'].startswith('inline')
        pdf = PdfReader(io.BytesIO(response.data))
        assert len(pdf.pages) == 2
        page = pdf.pages[0]
        assert float(page.mediabox.width) == 288 and float(page.mediabox.height) == 432
        label = page.extract_text()
        assert 'Acme Logistics' in label and 'DELIVER TO' in label and 'Receiver Road' in label and 'Sender Road' in label
        embedded_fonts = [font.get_object().get('/FontDescriptor') for font in page['/Resources']['/Font'].values()]
        assert sum(bool(descriptor and descriptor.get_object().get('/FontFile2')) for descriptor in embedded_fonts) == 2
        outputs.append(page.get_contents().get_data().count(b' re'))
    assert outputs[0] > 20 and outputs[1] == 0


def test_legacy_company_and_crm_upgrade_preserves_data_and_is_repeatable(tmp_path, monkeypatch):
    path = tmp_path/'legacy-admin.db'
    engine = create_engine(f'sqlite:///{path}')
    with engine.begin() as connection:
        connection.execute(text('CREATE TABLE lead (id INTEGER PRIMARY KEY, name VARCHAR(100) NOT NULL, email VARCHAR(120) NOT NULL, phone VARCHAR(30), subject VARCHAR(200), message TEXT NOT NULL, created_at DATETIME NOT NULL)'))
        connection.execute(text("INSERT INTO lead VALUES (1, 'Keep contact', 'keep@example.test', '9999999999', 'Old enquiry', 'Keep message', '2026-01-01 12:00:00')"))
    engine.dispose()
    monkeypatch.setenv('FLASK_ENV', 'development')
    upgraded = create_app({'INSTANCE_PATH': str(tmp_path/'legacy-instance'), 'SQLALCHEMY_DATABASE_URI': f'sqlite:///{path}', 'AUTO_CREATE_TABLES': True, 'RATELIMIT_ENABLED': False})
    for _ in range(2):
        assert upgraded.test_cli_runner().invoke(args=['upgrade-db']).exit_code == 0
    with upgraded.app_context():
        row = db.session.get(Lead, 1)
        assert row.name == 'Keep contact' and row.message == 'Keep message' and row.status == 'New'
        assert row.follow_up_date is None and row.notes is None
        assert Company.query.count() == 0
        db.session.remove()
        db.engine.dispose()


@pytest.mark.parametrize('path,method,data', [
    ('/admin/companies', 'post', {'name': 'Unauthenticated'}), ('/admin/companies/options', 'get', None),
    ('/admin/leads', 'post', {'name': 'Unauthenticated'}), ('/admin/leads/1', 'put', {}),
    ('/admin/leads/bulk', 'post', {'ids': [1], 'action': 'delete'}), ('/admin/leads/email-preview', 'post', {'ids': [1]}),
])
def test_company_and_crm_json_require_login(client, path, method, data):
    response = getattr(client, method)(path, json=data, headers={'Accept': 'application/json'})
    assert response.status_code == 401
