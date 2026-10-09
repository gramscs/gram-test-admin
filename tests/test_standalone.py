import base64
import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app import create_app
from app.models import Consignment, Lead, NewsletterSubscriber, db


def save(client, number="TEST001", **fields):
    response = client.post("/admin/consignments/save", json={"rows": [{
        "consignment_number": number, "status": "In Transit", **fields,
    }]})
    assert response.status_code == 200, response.get_json()
    assert response.get_json()["success"]
    return client.get("/admin/consignments/list", query_string={"search": number}).get_json()["rows"][0]


@pytest.mark.parametrize("path", ["/", "/admin", "/admin/dashboard", "/admin/consignments", "/admin/leads"])
def test_pages_require_login(client, path):
    response = client.get(path, follow_redirects=True)
    assert response.status_code == 200
    assert b"Admin Login" in response.data


def test_json_requests_require_login(client):
    response = client.get("/admin/consignments/list", headers={"Accept": "application/json"})
    assert response.status_code == 401
    assert response.get_json()["error"] == "Authentication required"


def test_invalid_login_and_logout(client, admin_client):
    response = client.get("/admin/logout", follow_redirects=True)
    assert b"Admin Login" in response.data
    response = client.post("/admin/login", data={"username": "admin", "password": "wrong"})
    assert b"Invalid username or password" in response.data
    assert client.get("/admin/dashboard").status_code == 302


@pytest.mark.parametrize("path,expected", [
    ("/admin/dashboard", b"Download Backup"),
    ("/admin/consignments", b"Internal Consignment Sheet"),
    ("/admin/leads", b"Customer Leads"),
])
def test_admin_pages_render_without_public_website(admin_client, path, expected):
    response = admin_client.get(path)
    assert response.status_code == 200
    assert expected in response.data
    assert b"https://cdn.jsdelivr.net" not in response.data


def test_all_browser_assets_are_bundled(client):
    for filename in [
        "vendor/bootstrap/bootstrap.min.css", "vendor/bootstrap/bootstrap.bundle.min.js",
        "vendor/space-grotesk/font.css", "vendor/space-grotesk/space-grotesk-latin-400-normal.woff2",
        "css/font-awesome.min.css", "fonts/fontawesome-webfont.woff2", "images/logo.png",
        "js/consignments.js", "js/admin/api.js", "js/admin/state.js", "js/admin/validation.js",
    ]:
        response = client.get("/static/" + filename)
        assert response.status_code == 200, filename
        assert len(response.data) > 0


def test_create_update_search_and_delete_shipment(admin_client, app):
    row = save(admin_client, pickup_address="New Delhi", drop_pincode="400001")
    changed = save(admin_client, id=row["id"], status="Delivered", drop_date="2026-01-01")
    assert changed["status"] == "Delivered"
    result = admin_client.get("/admin/consignments/list", query_string={"search": "Delivered", "per_page": 1}).get_json()
    assert result["total"] == 1
    response = admin_client.post("/admin/consignments/save", json={"rows": [], "deleted_ids": [row["id"]]})
    assert response.get_json()["deleted_count"] == 1
    with app.app_context():
        assert Consignment.query.count() == 0


def test_invalid_shipment_is_not_saved(admin_client):
    response = admin_client.post("/admin/consignments/save", json={"rows": [{"consignment_number": ""}]})
    assert response.status_code == 400
    assert admin_client.get("/admin/consignments/list").get_json()["total"] == 0


def test_delivery_proof_save_download_and_remove(admin_client, app):
    image = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jZQAAAABJRU5ErkJggg==")
    row = save(admin_client, pod_file_data="data:image/png;base64," + base64.b64encode(image).decode(), pod_file_name="proof.png", pod_file_type="image/png")
    pod_path = Path(app.instance_path) / "uploads" / row["pod_image"]
    assert pod_path.read_bytes() == image
    response = admin_client.get(f"/admin/consignments/{row['id']}/pod")
    assert response.status_code == 200
    assert response.data == image
    response = admin_client.delete(f"/admin/consignments/{row['id']}/pod")
    assert response.get_json()["success"]
    assert not pod_path.exists()
    assert admin_client.get(f"/admin/consignments/{row['id']}/pod").status_code == 404


def test_excel_import_duplicate_handling_and_exports(admin_client):
    template = admin_client.get("/admin/consignments/import-template.xlsx")
    assert template.status_code == 200
    workbook = load_workbook(io.BytesIO(template.data))
    workbook.active.append(["EXCEL001", "Delivered", "Pickup", "110017", "Origin", "2026-01-01", "Drop", "400001", "Destination", "2026-01-02"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()
    for expected in [b"Added: 1", b"Added: 0"]:
        response = admin_client.post("/admin/consignments/import", data={"file": (io.BytesIO(data), "import.xlsx")}, follow_redirects=True)
        assert response.status_code == 200
        assert expected in response.data
    exported = admin_client.get("/admin/consignments/export.xlsx")
    assert exported.status_code == 200
    sheet = load_workbook(io.BytesIO(exported.data)).active
    assert sheet.cell(2, 1).value == "EXCEL001"
    pdf = admin_client.get("/admin/consignments/export.pdf")
    assert pdf.status_code == 200
    assert pdf.data.startswith(b"%PDF")


def test_archive_option_and_endpoint_are_removed(admin_client, app):
    save(admin_client, "OLD001", status="Delivered", drop_date="2026-01-01")
    assert b"Archive Delivered" not in admin_client.get("/admin/consignments").data
    response = admin_client.post("/admin/consignments/archive", json={"before_date": "2026-06-01"})
    assert response.status_code == 404
    with app.app_context():
        assert Consignment.query.count() == 1


def test_leads_panel_rejection_and_json_backup(admin_client, app):
    with app.app_context():
        db.session.add_all([
            Lead(name="Valid enquiry", email="valid@example.com", phone="1234567890", message="Hello"),
            Lead(name="Missing phone", email="blank@example.com", phone="", message="Hello"),
            NewsletterSubscriber(email="subscriber@example.com"),
        ])
        db.session.commit()
    response = admin_client.get("/admin/leads")
    assert b"Valid enquiry" in response.data
    response = admin_client.post("/admin/leads/reject-empty-phone", follow_redirects=True)
    assert response.status_code == 200
    assert b"Missing phone" not in response.data
    with app.app_context():
        assert Lead.query.count() == 1
    payload = admin_client.get("/admin/generate-backup").get_json()
    assert payload["metadata"]["table_counts"]["leads"] == 1
    assert payload["leads"][0]["name"] == "Valid enquiry"
    assert payload["newsletter_subscribers"][0]["email"] == "subscriber@example.com"


def test_health_and_init_db_command(client, app):
    assert client.get("/health").get_json()["status"] == "ok"
    assert client.get("/health/db").get_json() == {"status": "ok", "database": "sqlite"}
    result = app.test_cli_runner().invoke(args=["init-db"])
    assert result.exit_code == 0
    assert "Database tables created" in result.output
    checked = app.test_cli_runner().invoke(args=["check-db"])
    assert checked.exit_code == 0
    assert "All required admin tables and columns are present" in checked.output


def test_database_persists_between_app_instances(admin_client, app):
    save(admin_client, "PERSIST001")
    restarted = create_app({
        "TESTING": True, "SECRET_KEY": "test-key", "INSTANCE_PATH": app.instance_path,
        "SQLALCHEMY_DATABASE_URI": app.config["SQLALCHEMY_DATABASE_URI"],
        "RATELIMIT_ENABLED": False, "AUTO_CREATE_TABLES": True,
    })
    with restarted.app_context():
        assert Consignment.query.count() == 1
        assert Consignment.query.first().consignment_number == "PERSIST001"
        db.session.remove()
        db.engine.dispose()


@pytest.mark.parametrize("missing", ["SECRET_KEY", "ADMIN_PASSWORD_HASH", "DATABASE_URL"])
def test_production_requires_configuration(app, monkeypatch, tmp_path, missing):
    monkeypatch.setenv("FLASK_ENV", "production")
    monkeypatch.setenv("SECRET_KEY", "test-only-key")
    monkeypatch.setenv("ADMIN_PASSWORD_HASH", "test-only-hash")
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'production.db'}")
    monkeypatch.delenv(missing)
    with pytest.raises(RuntimeError, match=missing):
        create_app({"INSTANCE_PATH": str(tmp_path / "production-instance")})


@pytest.mark.parametrize("scheme", ["postgres", "postgresql"])
def test_supabase_configuration_without_connecting(app, monkeypatch, tmp_path, scheme):
    monkeypatch.setenv("DATABASE_URL", f"{scheme}://postgres.example:dummy@aws-0-ap-south-1.pooler.supabase.com:5432/postgres")
    monkeypatch.delenv("AUTO_CREATE_TABLES", raising=False)
    application = create_app({"INSTANCE_PATH": str(tmp_path / "supabase-instance")})
    from sqlalchemy.engine import make_url
    uri = make_url(application.config["SQLALCHEMY_DATABASE_URI"])
    assert uri.get_backend_name() == "postgresql"
    assert uri.query["sslmode"] == "require"
    assert not application.config["AUTO_CREATE_TABLES"]
    assert application.config["SQLALCHEMY_ENGINE_OPTIONS"]["connect_args"]["connect_timeout"] == 10
    assert application.config["SQLALCHEMY_ENGINE_OPTIONS"]["pool_pre_ping"]
    with application.app_context():
        db.engine.dispose()


def test_supabase_preserves_certificate_verification_settings(app, monkeypatch, tmp_path):
    monkeypatch.setenv("DATABASE_URL", "postgresql://postgres:dummy@db.example.supabase.co:5432/postgres?sslmode=verify-full&sslrootcert=/tmp/test-root.crt")
    monkeypatch.delenv("AUTO_CREATE_TABLES", raising=False)
    application = create_app({"INSTANCE_PATH": str(tmp_path / "ssl-instance")})
    from sqlalchemy.engine import make_url
    query = make_url(application.config["SQLALCHEMY_DATABASE_URI"]).query
    assert query["sslmode"] == "verify-full"
    assert query["sslrootcert"] == "/tmp/test-root.crt"
    with application.app_context():
        db.engine.dispose()


def test_supabase_falls_back_to_sqlite_when_placeholder_is_left_in_local_dev(monkeypatch, tmp_path):
    monkeypatch.setenv("FLASK_ENV", "development")
    monkeypatch.setenv("DATABASE_URL", "postgresql://postgres.example:[YOUR-PASSWORD]@aws-0-ap-south-1.pooler.supabase.com:5432/postgres")
    application = create_app({"INSTANCE_PATH": str(tmp_path / "placeholder-instance")})
    assert application.config["SQLALCHEMY_DATABASE_URI"].startswith("sqlite:///")
    assert application.config["SQLALCHEMY_DATABASE_URI"].endswith("admin.db")
    with application.app_context():
        db.engine.dispose()


def test_supabase_rejects_unfilled_password_placeholder(app, monkeypatch, tmp_path):
    monkeypatch.setenv("FLASK_ENV", "production")
    monkeypatch.setenv("SECRET_KEY", "test-key")
    monkeypatch.setenv("ADMIN_PASSWORD_HASH", "test-hash")
    monkeypatch.setenv("DATABASE_URL", "postgresql://postgres.example:[YOUR-PASSWORD]@aws-0-ap-south-1.pooler.supabase.com:5432/postgres")
    with pytest.raises(RuntimeError, match="password placeholder"):
        create_app({"INSTANCE_PATH": str(tmp_path / "placeholder-instance")})


def test_check_db_reports_missing_structure_without_creating_it(app):
    with app.app_context():
        NewsletterSubscriber.__table__.drop(db.engine)
    result = app.test_cli_runner().invoke(args=["check-db"])
    assert result.exit_code == 1
    assert "table newsletter_subscriber" in result.output
    from sqlalchemy import inspect
    with app.app_context():
        assert not inspect(db.engine).has_table("newsletter_subscriber")


def test_check_db_does_not_expose_connection_errors(app, monkeypatch):
    def connection_failure(*args, **kwargs):
        raise RuntimeError("fake-sensitive-connection-detail")
    monkeypatch.setattr(db.session, "execute", connection_failure)
    result = app.test_cli_runner().invoke(args=["check-db"])
    assert result.exit_code == 1
    assert "Database connection failed" in result.output
    assert "fake-sensitive-connection-detail" not in result.output
