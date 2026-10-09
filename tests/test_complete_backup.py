import hashlib
import io
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfWriter
from sqlalchemy import Column, Integer, String, Table

from app.models import Company, CompanyLocation, Consignment, Lead, NewsletterSubscriber, db


def image_bytes():
    buffer = io.BytesIO()
    Image.new('RGB', (6, 6), 'green').save(buffer, format='PNG')
    return buffer.getvalue()


def pdf_bytes():
    writer = PdfWriter()
    writer.add_blank_page(width=288, height=432)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def unpack(response):
    assert response.status_code == 200, response.get_data(as_text=True) if response.mimetype != 'application/zip' else ''
    assert response.mimetype == 'application/zip'
    archive = ZipFile(io.BytesIO(response.data))
    assert archive.testzip() is None
    return archive, json.loads(archive.read('data.json')), json.loads(archive.read('document_report.json'))


def test_complete_backup_contains_all_tables_fields_and_documents_in_data_order(admin_client, app):
    with app.app_context():
        master = Company(name='Client master', address='Registered address', active=False)
        master.locations = [CompanyLocation(kind='pickup', label='Warehouse', address='Pickup address', is_default=True), CompanyLocation(kind='drop', label='Office', address='Drop address', is_default=True)]
        db.session.add(master)
        first = Consignment(consignment_number='ZZZ001', company=master, eta_debug_json='{"keep":"all fields"}', pod_image='proof.png', pod_original_name='signed proof.png', invoice_file='invoice.pdf', invoice_original_name='bill.pdf')
        second = Consignment(consignment_number='AAA001', company=master, pod_image='other.png', pod_original_name='signed proof.png')
        db.session.add_all([first, second, Lead(name='Legacy customer', email='old@example.test', message='=KEEP()', notes='Keep notes', company_name='Old company'), NewsletterSubscriber(email='old-subscriber@example.test')])
        db.session.commit()
        before = {table.name: db.session.query(table).count() for table in db.metadata.sorted_tables}
    folder = Path(app.instance_path, 'uploads')
    folder.mkdir()
    content = {'proof.png': image_bytes(), 'other.png': image_bytes(), 'invoice.pdf': pdf_bytes(), 'old-unlinked.png': image_bytes()}
    for name, data in content.items():
        (folder/name).write_bytes(data)
    response = admin_client.get('/admin/generate-backup')
    archive, data, report = unpack(response)
    assert response.headers['X-Backup-Status'] == 'complete'
    assert response.headers['Cache-Control'] == 'private, no-store'
    assert data['metadata']['table_counts'] == {'companies': 1, 'company_locations': 2, 'consignments': 2, 'leads': 1, 'newsletter_subscribers': 1, 'mis_view': 0, 'mis_report': 0}
    assert data['companies'][0]['active'] is False
    assert data['company_locations'][0]['address'] == 'Pickup address'
    assert data['leads'][0]['notes'] == 'Keep notes'
    assert data['newsletter_subscribers'][0]['email'] == 'old-subscriber@example.test'
    assert data['consignments'][0]['eta_debug_json'] == '{"keep":"all fields"}'
    assert [row['consignment_number'] for row in data['consignments']] == ['ZZZ001', 'AAA001']
    assert [row['backup_order'] for row in data['consignments']] == [1, 2]
    first, second = data['consignments']
    assert '/000001_ZZZ001_' in first['pod_backup_path'] and '/000002_AAA001_' in second['pod_backup_path']
    assert first['pod_backup_path'] != second['pod_backup_path']
    assert archive.read(first['pod_backup_path']) == content['proof.png']
    assert archive.read(first['invoice_backup_path']) == content['invoice.pdf']
    assert archive.read(second['pod_backup_path']) == content['other.png']
    entries = [entry for entry in report['documents'] if entry['status'] == 'included']
    assert len(entries) == 4
    for entry in entries:
        assert entry['sha256'] == hashlib.sha256(archive.read(entry['archive_path'])).hexdigest()
    assert any(entry['kind'] == 'unlinked_upload' for entry in entries)
    workbook = load_workbook(io.BytesIO(archive.read('admin-data.xlsx')))
    assert set(workbook.sheetnames) == {'companies', 'company_locations', 'consignments', 'leads', 'newsletter_subscribers', 'mis_view', 'mis_report', 'document_report', 'inventory_errors'}
    sheet = workbook['consignments']
    values = list(sheet.values)
    rows = [dict(zip(values[0], row)) for row in values[1:]]
    assert [row['backup_order'] for row in rows] == [1, 2]
    assert rows[0]['pod_backup_path'] == first['pod_backup_path']
    lead_sheet = workbook['leads']
    headers = [cell.value for cell in lead_sheet[1]]
    message = lead_sheet.cell(2, headers.index('message')+1)
    assert message.value == '=KEEP()' and message.data_type == 's'
    with app.app_context():
        assert {table.name: db.session.query(table).count() for table in db.metadata.sorted_tables} == before
    assert set(path.name for path in folder.iterdir()) == set(content)


def test_unlinked_nested_uploads_and_missing_files_are_reported(admin_client, app):
    with app.app_context():
        db.session.add(Consignment(consignment_number='MISSING', pod_image='gone.png'))
        db.session.commit()
    root = Path(app.instance_path, 'uploads')
    (root/'legacy').mkdir(parents=True)
    (root/'legacy'/'proof.png').write_bytes(image_bytes())
    (root/'proof.png').write_bytes(b'old original document')
    archive, data, report = unpack(admin_client.get('/admin/generate-backup'))
    assert data['metadata']['backup_status'] == 'incomplete'
    assert report['summary']['document_counts'] == {'missing': 1, 'not_attached': 1, 'included': 2}
    assert report['documents'][0]['archive_path'] is None and report['documents'][0]['status'] == 'missing'
    paths = [row['archive_path'] for row in report['documents'] if row['status'] == 'included']
    assert len(set(paths)) == 2 and all(path.startswith('other_uploads/') for path in paths)
    assert b'INCOMPLETE' in archive.read('README.txt')
    response = admin_client.get('/admin/generate-backup')
    assert '_incomplete.zip' in response.headers['Content-Disposition']


def test_snapshot_automatically_includes_additional_app_owned_table(admin_client, app):
    with app.app_context():
        table = Table('additional_admin_records', db.metadata, Column('id', Integer, primary_key=True), Column('details', String(100)))
        try:
            table.create(db.engine)
            db.session.execute(table.insert().values(id=1, details='Keep all admin data'))
            db.session.commit()
            archive, data, _ = unpack(admin_client.get('/admin/generate-backup'))
            assert data['additional_admin_records'] == [{'id': 1, 'details': 'Keep all admin data'}]
            assert 'additional_admin_records' in load_workbook(io.BytesIO(archive.read('admin-data.xlsx'))).sheetnames
        finally:
            table.drop(db.engine)
            db.metadata.remove(table)


def test_backup_rejects_escaped_paths_external_urls_and_outside_symlinks(admin_client, app):
    outside = Path(app.instance_path, 'outside.png')
    outside.write_bytes(b'private-outside-content')
    root = Path(app.instance_path, 'uploads')
    root.mkdir()
    (root/'outside-link.png').symlink_to(outside)
    with app.app_context():
        db.session.add(Consignment(consignment_number='../UNSAFE?', pod_image='../outside.png', invoice_file='https://example.test/private.pdf'))
        db.session.commit()
    archive, data, report = unpack(admin_client.get('/admin/generate-backup'))
    assert data['metadata']['backup_status'] == 'incomplete'
    assert all(entry['status'] == 'invalid_reference' for entry in report['documents'] if entry.get('source_reference'))
    assert all('..' not in name.split('/') and not name.startswith('/') for name in archive.namelist())
    assert not any(name.startswith(('documents/', 'other_uploads/')) for name in archive.namelist())
    assert outside.read_bytes() == b'private-outside-content'


def test_private_supabase_backup_lists_unlinked_uploads_and_nested_folders(admin_client, app, monkeypatch):
    from app.admin import consignment_controller
    objects = {'consignments/proof.png': image_bytes(), 'consignments/unlinked.pdf': pdf_bytes(), 'consignments/legacy/old.png': image_bytes()}
    with app.app_context():
        db.session.add(Consignment(consignment_number='REMOTE', pod_image='supabase:private-uploads/consignments/proof.png'))
        db.session.commit()
    class Bucket:
        def list(self, folder, options):
            if folder == 'consignments':
                return [{'name': 'legacy', 'id': None, 'metadata': None}, {'name': 'proof.png', 'id': '1'}, {'name': 'unlinked.pdf', 'id': '2'}]
            return [{'name': 'old.png', 'id': '3'}]
        def download(self, path):
            return objects[path]
    class Storage:
        def from_(self, bucket):
            assert bucket == 'private-uploads'
            return Bucket()
    class Client:
        storage = Storage()
    monkeypatch.setattr(consignment_controller, '_get_supabase_client', lambda: Client())
    archive, _, report = unpack(admin_client.get('/admin/generate-backup'))
    assert report['summary']['status'] == 'complete'
    entries = [entry for entry in report['documents'] if entry['status'] == 'included']
    assert len(entries) == 3
    assert sum(entry['kind'] == 'unlinked_upload' for entry in entries) == 2
    for entry in entries:
        reference = entry['source_reference'].split('private-uploads/', 1)[1]
        assert archive.read(entry['archive_path']) == objects[reference]


def test_unavailable_supabase_inventory_and_download_are_explicit(admin_client, app, monkeypatch):
    from app.admin import consignment_controller
    with app.app_context():
        db.session.add(Consignment(consignment_number='UNAVAILABLE', pod_image='supabase:private-uploads/consignments/proof.png'))
        db.session.commit()
    class Bucket:
        def list(self, *args):
            raise Exception('storage SDK error with sensitive details')
        def download(self, *args):
            raise Exception('storage SDK error with sensitive details')
    class Storage:
        def from_(self, *args):
            return Bucket()
    class Client:
        storage = Storage()
    monkeypatch.setattr(consignment_controller, '_get_supabase_client', lambda: Client())
    archive, data, report = unpack(admin_client.get('/admin/generate-backup'))
    assert report['summary']['status'] == 'incomplete'
    assert report['documents'][0]['status'] == 'unavailable'
    assert report['inventory_errors'][0]['source'] == 'supabase_uploads'
    assert data['consignments'][0]['consignment_number'] == 'UNAVAILABLE'
    assert b'sensitive details' not in archive.read('document_report.json')


def test_json_only_contains_all_fields_and_does_not_read_uploads(admin_client, app, monkeypatch):
    with app.app_context():
        db.session.add(Consignment(consignment_number='DATA', eta_debug_json='Keep all fields', pod_image='missing.png'))
        db.session.commit()
    from app.admin import backup
    monkeypatch.setattr(backup, 'build_complete_backup', lambda *args: pytest.fail('JSON must not retrieve files'))
    response = admin_client.get('/admin/generate-backup?format=json')
    assert response.mimetype == 'application/json'
    assert response.get_json()['consignments'][0]['eta_debug_json'] == 'Keep all fields'


def test_empty_backup_contains_all_table_sheets_and_report(admin_client):
    archive, data, report = unpack(admin_client.get('/admin/generate-backup'))
    assert data['metadata']['total_rows'] == 0
    assert report['summary']['status'] == 'complete'
    assert load_workbook(io.BytesIO(archive.read('admin-data.xlsx')))['companies'].cell(1, 2).value == 'name'


def test_complete_backup_requires_login(client):
    response = client.get('/admin/generate-backup', headers={'Accept': 'application/json'})
    assert response.status_code == 401


def test_invalid_format_and_database_failure_return_no_archive(admin_client, app):
    assert admin_client.get('/admin/generate-backup?format=unknown').status_code == 400
    with app.app_context():
        CompanyLocation.__table__.drop(db.engine)
    response = admin_client.get('/admin/generate-backup')
    assert response.status_code == 500 and response.mimetype == 'application/json'


def test_supabase_inventory_includes_files_beyond_first_page(admin_client, app, monkeypatch):
    from app.admin import consignment_controller
    names = [f'file-{index:03d}.bin' for index in range(105)]
    calls = []
    with app.app_context():
        db.session.add(Consignment(consignment_number='PAGED', pod_image='supabase:private-uploads/consignments/file-000.bin'))
        db.session.commit()
    class Bucket:
        def list(self, folder, options):
            calls.append(options['offset'])
            return [{'name': name, 'id': name} for name in names[options['offset']:options['offset']+options['limit']]]
        def download(self, path):
            return path.encode()
    class Storage:
        def from_(self, bucket):
            return Bucket()
    class Client:
        storage = Storage()
    monkeypatch.setattr(consignment_controller, '_get_supabase_client', lambda: Client())
    archive, _, report = unpack(admin_client.get('/admin/generate-backup'))
    assert calls == [0, 100]
    assert report['summary']['document_counts']['included'] == 105
    assert report['summary']['status'] == 'complete'
    assert len([name for name in archive.namelist() if name.startswith('other_uploads/')]) == 104


def test_unreadable_local_upload_directory_is_not_silently_marked_complete(admin_client, app, monkeypatch):
    from app.admin import backup
    root = Path(app.instance_path, 'uploads')
    root.mkdir()
    (root/'readable.png').write_bytes(image_bytes())
    def listing(path, onerror, followlinks):
        assert not followlinks
        onerror(PermissionError('cannot list nested folder'))
        yield str(root), [], ['readable.png']
    monkeypatch.setattr(backup.os, 'walk', listing)
    archive, _, report = unpack(admin_client.get('/admin/generate-backup'))
    assert report['summary']['status'] == 'incomplete'
    assert report['summary']['document_counts']['included'] == 1
    assert report['inventory_errors'][0]['source'] == 'local_uploads'
    assert archive.read(report['documents'][0]['archive_path']) == image_bytes()
