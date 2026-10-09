import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import Column, Integer, MetaData, String, Table, create_engine

from app import create_app
from app.models import Consignment, db


def saved(client, **values):
    payload = {"consignment_number": "AWB12345678901234567890", "status": "In Transit", "identifier_type": "AWB", "pieces": 3, "chargeable_weight": "12.345", "chargeable_volume": "0.125", "pickup_address": "12 Sender Road, Delhi", "drop_address": "34 Receiver Road, Mumbai", "pickup_pincode": "110001", "drop_pincode": "400001", **values}
    response = client.post("/admin/consignments/save", json={"rows": [payload]})
    assert response.status_code == 200, response.get_json()
    return client.get("/admin/consignments/list").get_json()["rows"][0]


@pytest.mark.parametrize("identifier_type", ["LRN", "Order ID", "AWB"])
def test_shipment_fields_save_reload_and_legacy_update(admin_client, identifier_type):
    row = saved(admin_client, identifier_type=identifier_type)
    assert row["identifier_type"] == identifier_type
    assert row["pieces"] == 3
    assert Decimal(row["chargeable_weight"]) == Decimal("12.345")
    assert Decimal(row["chargeable_volume"]) == Decimal("0.125")
    response = admin_client.post("/admin/consignments/save", json={"rows": [{"id": row["id"], "consignment_number": row["consignment_number"], "status": "Delivered"}]})
    assert response.status_code == 200
    updated = admin_client.get("/admin/consignments/list").get_json()["rows"][0]
    assert updated["pieces"] == 3
    assert updated["identifier_type"] == identifier_type
    assert Decimal(updated["chargeable_weight"]) == Decimal("12.345")


@pytest.mark.parametrize("values", [
    {"identifier_type": "Invalid"}, {"pieces": 0}, {"pieces": -1}, {"pieces": "1.5"},
    {"pieces": 10001}, {"chargeable_weight": "-0.1"}, {"chargeable_volume": "NaN"},
    {"chargeable_weight": "Infinity"}, {"chargeable_volume": "0.0001"},
    {"chargeable_weight": "1000000000"}, {"consignment_number": "X" * 65},
])
def test_invalid_shipment_fields_do_not_persist(admin_client, app, values):
    response = admin_client.post("/admin/consignments/save", json={"rows": [{"consignment_number": "INVALID01", "status": "In Transit", **values}]})
    assert response.status_code == 400
    with app.app_context():
        assert Consignment.query.count() == 0


def test_weight_and_volume_can_be_zero_or_cleared(admin_client):
    row = saved(admin_client, chargeable_weight="0", chargeable_volume="0")
    assert Decimal(row["chargeable_weight"]) == 0
    response = admin_client.post("/admin/consignments/save", json={"rows": [{"id": row["id"], "consignment_number": row["consignment_number"], "chargeable_weight": "", "chargeable_volume": ""}]})
    assert response.status_code == 200
    updated = admin_client.get("/admin/consignments/list").get_json()["rows"][0]
    assert updated["chargeable_weight"] is None and updated["chargeable_volume"] is None


def test_new_fields_excel_roundtrip(admin_client):
    saved(admin_client, identifier_type="Order ID")
    response = admin_client.get("/admin/consignments/export.xlsx")
    sheet = load_workbook(io.BytesIO(response.data)).active
    values = list(sheet.iter_rows(values_only=True))
    exported = dict(zip(values[0], values[1]))
    assert exported["identifier_type"] == "Order ID"
    assert exported["pieces"] == 3
    assert exported["chargeable_weight"] == 12.345
    template = load_workbook(io.BytesIO(admin_client.get("/admin/consignments/import-template.xlsx").data))
    headers = [cell.value for cell in template.active[1]]
    row = {"consignment_number": "EXCELNEW01", "status": "In Transit", "identifier_type": "LRN", "pieces": 2, "chargeable_weight": 5.25, "chargeable_volume": 0.4}
    template.active.append([row.get(header) for header in headers])
    buffer = io.BytesIO()
    template.save(buffer)
    result = admin_client.post("/admin/consignments/import", data={"file": (io.BytesIO(buffer.getvalue()), "shipments.xlsx")}, follow_redirects=True)
    assert b"Added: 1" in result.data
    imported = admin_client.get("/admin/consignments/list?search=EXCELNEW01").get_json()["rows"][0]
    assert imported["pieces"] == 2
    assert Decimal(imported["chargeable_weight"]) == Decimal("5.25")


def test_labels_pdf_size_piece_count_and_contents(admin_client):
    row = saved(admin_client)
    assert b"B2B Shipping Labels" in admin_client.get("/admin/labels").data
    response = admin_client.post("/admin/labels/generate", data={"consignment_ids": [row["id"]]})
    assert response.status_code == 200
    assert response.mimetype == "application/pdf"
    pdf = PdfReader(io.BytesIO(response.data))
    assert len(pdf.pages) == 3
    for index, page in enumerate(pdf.pages, 1):
        assert float(page.mediabox.width) == 288 and float(page.mediabox.height) == 432
        text = page.extract_text()
        assert f"Piece {index} of 3" in text
        assert row["consignment_number"] in text
        assert "12 Sender Road" in text and "34 Receiver Road" in text
        assert "12.345 kg" in text and "0.125 m3" in text
        assert b" re" in page.get_contents().get_data()  # Barcode bars are drawn as rectangles.


def test_label_errors_and_duplicate_selection(admin_client):
    assert admin_client.post("/admin/labels/generate").status_code == 400
    assert admin_client.post("/admin/labels/generate", data={"consignment_ids": [999]}).status_code == 400
    row = saved(admin_client, pieces=501)
    assert admin_client.post("/admin/labels/generate", data={"consignment_ids": [row["id"]]}).status_code == 400
    row = saved(admin_client, id=row["id"], pieces=1)
    result = admin_client.post("/admin/labels/generate", data={"consignment_ids": [row["id"], row["id"]]})
    assert len(PdfReader(io.BytesIO(result.data)).pages) == 1


def test_labels_require_login(client):
    assert client.get("/admin/labels").status_code == 302
    assert client.post("/admin/labels/generate", data={"consignment_ids": [1]}).status_code == 302


def test_old_sqlite_database_upgrades_without_losing_records(app, tmp_path):
    path = tmp_path / "legacy.db"
    engine = create_engine(f"sqlite:///{path}")
    metadata = MetaData()
    new_fields = {"identifier_type", "pieces", "chargeable_weight", "chargeable_volume"}
    columns = [Column(column.name, String(16) if column.name == "consignment_number" else column.type, primary_key=column.primary_key) for column in Consignment.__table__.columns if column.name not in new_fields]
    old = Table("consignment", metadata, *columns)
    metadata.create_all(engine)
    with engine.begin() as connection:
        connection.execute(old.insert().values(id=1, consignment_number="KEEP001", status="Delivered", pickup_address="Keep existing address"))
    engine.dispose()
    upgraded = create_app({"INSTANCE_PATH": str(tmp_path / "legacy-instance"), "SQLALCHEMY_DATABASE_URI": f"sqlite:///{path}", "AUTO_CREATE_TABLES": True, "RATELIMIT_ENABLED": False})
    with upgraded.app_context():
        row = db.session.get(Consignment, 1)
        assert row.consignment_number == "KEEP001" and row.pickup_address == "Keep existing address"
        assert row.identifier_type == "LRN" and row.pieces == 1
        assert row.chargeable_weight is None
    for _ in range(2):
        assert upgraded.test_cli_runner().invoke(args=["upgrade-db"]).exit_code == 0
    with upgraded.app_context():
        assert Consignment.query.count() == 1
        db.session.remove()
        db.engine.dispose()
