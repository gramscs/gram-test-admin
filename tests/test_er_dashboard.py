"""Operate the dashboard against a fresh migrated database, without legacy tables."""
import base64
import io
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from sqlalchemy import inspect, select
from werkzeug.security import generate_password_hash

from app import create_app
from app.models import db
from database.models import Base
from migrations.shipment_v1 import RECORDS, migrate
from test_shipment_migration import populated, legacy_engine
from test_complete_database_import import disposable_engine


@pytest.fixture
def migrated_admin(populated, tmp_path, monkeypatch):
    source, storage, content = populated
    with disposable_engine(source.dialect.name, tmp_path / 'er-dashboard.db') as destination:
        migrate(destination, storage, source_engine=source, fresh_target=True)
        monkeypatch.setenv('FLASK_ENV', 'development')
        monkeypatch.setenv('DATABASE_URL', destination.url.render_as_string(hide_password=False))
        monkeypatch.setenv('ADMIN_USERNAME', 'admin')
        monkeypatch.delenv('SUPABASE_URL', raising=False)
        monkeypatch.delenv('SUPABASE_KEY', raising=False)
        application = create_app({'TESTING': True, 'SECRET_KEY': 'er-test-key',
            'INSTANCE_PATH': str(storage.root.parent), 'AUTO_CREATE_TABLES': False,
            'RATELIMIT_ENABLED': False, 'ADMIN_DATABASE_MODEL': 'auto'})
        from app.admin import auth
        monkeypatch.setattr(auth, 'ADMIN_PASSWORD_HASH', generate_password_hash('test-password'))
        result = application.test_cli_runner().invoke(args=['upgrade-db'])
        assert result.exit_code == 0, result.output
        assert 'ER backend ready' in result.output
        assert 'consignment' not in inspect(destination).get_table_names()
        assert 'company' not in inspect(destination).get_table_names()
        client = application.test_client()
        assert client.post('/admin/login', data={'username': 'admin', 'password': 'test-password'}).status_code == 302
        yield application, client, storage, destination
        with application.app_context():
            db.session.remove()
            db.engine.dispose()


@pytest.mark.parametrize('path', ['/admin/dashboard', '/admin/consignments', '/admin/companies', '/admin/labels', '/admin/mis', '/admin/companies/options', '/admin/consignments/list'])
def test_pages_use_the_migrated_model(migrated_admin, path):
    app, client, storage, destination = migrated_admin
    response = client.get(path)
    assert response.status_code == 200, response.get_data(as_text=True)
    if path == '/admin/consignments/list':
        rows = response.get_json()['rows']
        assert len(rows) == 2
        row = next(row for row in rows if row['consignment_number'] == 'LRN-001')
        assert row['pickup_pincode'] == '110001' and row['pickup_tag'] == 'Warehouse A'
        assert row['pod_image'] == 'local:pod.png'
        assert row['company_id'] and '-' in row['id']
        assert next(row for row in rows if row['consignment_number'] == 'AWB-002')['pickup_date'] == 'next Tuesday'


def test_er_check_db_and_upgrade_are_repeatable(migrated_admin):
    app, client, storage, destination = migrated_admin
    result = app.test_cli_runner().invoke(args=['check-db'])
    assert result.exit_code == 0, result.output
    assert 'All required admin tables and columns are present' in result.output
    result = app.test_cli_runner().invoke(args=['upgrade-db'])
    assert result.exit_code == 0, result.output
    assert client.get('/admin/consignments/list').get_json()['total'] == 2


def test_uuid_shipment_updates_extras_and_exports(migrated_admin):
    app, client, storage, destination = migrated_admin
    original = next(row for row in client.get('/admin/consignments/list').get_json()['rows'] if row['consignment_number'] == 'LRN-001')
    changed = {**original, 'pickup_tag': 'Updated warehouse', 'pickup_pincode': '110002', 'status': 'Delivered'}
    response = client.post('/admin/consignments/save', json={'rows': [changed]})
    assert response.status_code == 200, response.get_json()
    response = client.get('/admin/consignments/list', query_string={'search': 'Updated warehouse'})
    assert response.get_json()['total'] == 1
    for path in ('/admin/consignments/export.xlsx', '/admin/consignments/export.pdf'):
        exported = client.get(path)
        assert exported.status_code == 200, exported.get_json()
    with destination.begin() as connection:
        archived = connection.execute(select(RECORDS).where(RECORDS.c.source_table == 'consignment', RECORDS.c.source_id == '11')).mappings().one()
        assert archived['source_values']['pickup_tag'] == 'Warehouse A'
    deleted = client.post('/admin/consignments/save', json={'rows': [], 'deleted_ids': [original['id']]})
    assert deleted.status_code == 200 and deleted.get_json()['deleted_count'] == 1


def test_uuid_companies_and_label_printing(migrated_admin):
    app, client, storage, destination = migrated_admin
    company = client.get('/admin/companies/options').get_json()['companies'][0]
    assert '-' in company['id']
    update = client.put('/admin/companies/' + company['id'], json={**company, 'address': 'Updated registered office'})
    assert update.status_code == 200, update.get_json()
    shipment = client.get('/admin/consignments/list').get_json()['rows'][0]
    printed = client.post('/admin/labels/generate', data={'consignment_ids': [shipment['id']]})
    assert printed.status_code == 200 and printed.data.startswith(b'%PDF'), printed.get_json()


def test_uuid_documents_and_saved_mis(migrated_admin):
    app, client, storage, destination = migrated_admin
    shipment = next(row for row in client.get('/admin/consignments/list').get_json()['rows'] if row['consignment_number'] == 'LRN-001')
    proof = client.get(f"/admin/consignments/{shipment['id']}/pod")
    assert proof.status_code == 200
    image = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jZQAAAABJRU5ErkJggg==')
    upload = client.post(f"/admin/consignments/{shipment['id']}/pod", data={'file': (io.BytesIO(image), 'new-proof.png')})
    assert upload.status_code == 200, upload.get_json()
    with destination.begin() as connection:
        documents = connection.execute(select(Base.metadata.tables['shipment_documents']).where(Base.metadata.tables['shipment_documents'].c.document_type == 'POD')).mappings().all()
        assert len(documents) == 3
        assert sum(doc['is_current'] for doc in documents) == 2
    assert (storage.root / 'pod.png').exists()
    view = client.post('/admin/mis/views', json={'name': 'Fresh view', 'format': 'csv', 'filters': {'period': 'all', 'company': shipment['company_id']}, 'columns': ['consignment_number', 'pickup_tag']})
    assert view.status_code == 200, view.get_json()
    generated = client.post('/admin/mis/reports', json={'view_id': view.get_json()['view']['id'], 'name': 'Fresh report', 'format': 'pdf', 'filters': {'period': 'all', 'company': shipment['company_id']}, 'columns': ['consignment_number', 'pickup_tag']})
    assert generated.status_code == 200, generated.get_json()
    download = client.get(generated.get_json()['download_url'])
    assert download.status_code == 200 and download.data.startswith(b'%PDF')
    with destination.begin() as connection:
        from uuid import UUID
        reports = Base.metadata.tables['mis_reports']
        saved = connection.execute(select(reports).where(reports.c.id == UUID(generated.get_json()['id']))).mappings().one()
        assert str(saved['view_id']) == view.get_json()['view']['id']


def test_er_backup_contains_model_archive_support_and_all_files(migrated_admin):
    app, client, storage, destination = migrated_admin
    result = client.get('/admin/generate-backup?format=json')
    assert result.status_code == 200, result.get_json()
    data = result.get_json()
    assert set(Base.metadata.tables) <= set(data)
    assert 'admin_record_extras' in data and 'admin_migration_records' in data
    assert len(data['consignments']) == 2
    assert len(data['admin_migration_records']) == 9
    assert isinstance(data['mis_reports'][0]['metrics_snapshot'], dict)
    result = client.get('/admin/generate-backup')
    assert result.status_code == 200, result.get_json()
    with ZipFile(io.BytesIO(result.data)) as archive:
        assert archive.testzip() is None
        data = json.loads(archive.read('data.json'))
        report = json.loads(archive.read('document_report.json'))
        assert report['summary']['status'] == 'complete'
        assert len(data['admin_migration_records']) == 9
        assert 'admin-data.xlsx' in archive.namelist()
        assert any(name.startswith('mis_reports/') for name in archive.namelist())
        assert any(name.startswith('file_history/') for name in archive.namelist())


def test_er_create_history_attribution_and_preset_company_change(migrated_admin):
    app, client, storage, destination = migrated_admin
    company = client.get('/admin/companies/options').get_json()['companies'][0]
    created = client.post('/admin/consignments/save', json={'rows': [{'consignment_number': 'UUID-NEW', 'status': 'Pickup Scheduled', 'company_id': company['id'], 'pickup_tag': 'New pickup', 'pieces': 3}]})
    assert created.status_code == 200, created.get_json()
    row = client.get('/admin/consignments/list', query_string={'search': 'UUID-NEW'}).get_json()['rows'][0]
    changed = client.post('/admin/consignments/save', json={'rows': [{**row, 'status': 'In Transit'}]})
    assert changed.status_code == 200, changed.get_json()
    with destination.begin() as connection:
        events = connection.execute(select(Base.metadata.tables['shipment_events'])).mappings().all()
        assert [event['status'] for event in events] == ['Pickup Scheduled', 'In Transit']
        assert all(event['recorded_by'] is not None for event in events)
        logs = connection.execute(select(Base.metadata.tables['audit_logs'])).mappings().all()
        assert any(log['action'] == 'update' and log['entity_type'] == 'consignments' for log in logs)
    deleted = client.post('/admin/consignments/save', json={'rows': [], 'deleted_ids': [row['id']]})
    assert deleted.status_code == 200
    with destination.begin() as connection:
        assert not connection.execute(select(Base.metadata.tables['shipment_events'])).first()


def test_er_report_rename_preserves_snapshot_and_delete_cleans_its_file(migrated_admin):
    app, client, storage, destination = migrated_admin
    generated = client.post('/admin/mis/reports', json={'name': 'Delete report', 'format': 'csv', 'filters': {'period': 'all'}, 'columns': ['consignment_number']})
    assert generated.status_code == 200, generated.get_json()
    id = generated.get_json()['id']
    from uuid import UUID
    with destination.begin() as connection:
        table = Base.metadata.tables['mis_reports']
        original = dict(connection.execute(select(table).where(table.c.id == UUID(id))).mappings().one())
    renamed = client.put('/admin/mis/reports/' + id, json={'name': 'Renamed report', 'notes': 'Added note'})
    assert renamed.status_code == 200, renamed.get_json()
    with destination.begin() as connection:
        after = dict(connection.execute(select(table).where(table.c.id == UUID(id))).mappings().one())
        assert after == {**original, 'name': 'Renamed report'}
        files = Base.metadata.tables['files']
        file = connection.execute(select(files).where(files.c.id == original['file_id'])).mappings().one()
        path = storage.root / file['storage_path'].removeprefix('local:')
        assert path.exists()
    removed = client.delete('/admin/mis/reports/' + id, json={})
    assert removed.status_code == 200, removed.get_json()
    assert not path.exists()
    assert client.get(generated.get_json()['download_url']).status_code == 404


def test_staged_negative_ids_save_as_new_uuid_records(migrated_admin):
    app, client, storage, destination = migrated_admin
    response = client.post('/admin/consignments/save', json={'rows': [
        {'id': '-1', 'consignment_number': 'STAGED-001', 'status': 'In Transit'},
        {'id': -2, 'consignment_number': 'STAGED-002', 'pieces': 2}]})
    assert response.status_code == 200, response.get_json()
    rows = client.get('/admin/consignments/list', query_string={'search': 'STAGED-'}).get_json()['rows']
    assert len(rows) == 2 and all('-' in row['id'] and not row['id'].startswith('-') for row in rows)


def test_er_authorization_checks_active_identity_and_viewer_writes(migrated_admin):
    app, client, storage, destination = migrated_admin
    from sqlalchemy import update
    users = Base.metadata.tables['admin_users']
    with destination.begin() as connection:
        connection.execute(update(users).where(users.c.username == 'admin').values(role='viewer'))
    assert client.get('/admin/consignments/list').status_code == 200
    assert client.post('/admin/consignments/save', json={'rows': []}).status_code == 403
    with destination.begin() as connection:
        connection.execute(update(users).where(users.c.username == 'admin').values(active=False))
    assert client.get('/admin/consignments/list').status_code == 401
    assert client.post('/admin/login', data={'username': 'admin', 'password': 'test-password'}).status_code == 200


def test_er_reimport_restores_latest_extras_without_mutating_archives(migrated_admin, tmp_path):
    app, client, storage, source = migrated_admin
    row = client.get('/admin/consignments/list').get_json()['rows'][0]
    result = client.post('/admin/consignments/save', json={'rows': [{**row, 'pickup_tag': 'Edited after first migration'}]})
    assert result.status_code == 200, result.get_json()
    from app.er_adapter import RecordExtras, prepare_backend
    with disposable_engine(source.dialect.name, tmp_path / 'second-import.db') as target:
        migrate(target, storage, source_engine=source, source_format='er', fresh_target=True)
        prepare_backend(target)
        with target.begin() as connection:
            from uuid import UUID
            extra = connection.execute(select(RecordExtras.__table__).where(
                RecordExtras.entity_type == 'consignments', RecordExtras.entity_id == UUID(row['id']))).mappings().one()
            assert extra['data']['pickup_tag'] == 'Edited after first migration'


def test_er_demo_seed_does_not_recreate_legacy_tables(migrated_admin):
    app, client, storage, destination = migrated_admin
    if destination.dialect.name != 'sqlite':
        pytest.skip('Demo seeding intentionally supports local SQLite only.')
    result = app.test_cli_runner().invoke(args=['seed-demo'])
    assert result.exit_code == 0, result.output
    assert 'consignment' not in inspect(destination).get_table_names()
    result = app.test_cli_runner().invoke(args=['seed-demo'])
    assert result.exit_code == 0 and 'Added 0 demo' in result.output
    assert client.get('/admin/consignments/list').get_json()['total'] == 27


def test_er_excel_import_keeps_typed_dates_and_skips_duplicates(migrated_admin):
    app, client, storage, destination = migrated_admin
    from openpyxl import load_workbook
    template = client.get('/admin/consignments/import-template.xlsx')
    assert template.status_code == 200
    workbook = load_workbook(io.BytesIO(template.data))
    workbook.active.append(['ER-EXCEL-001', 'Delivered', 'Pickup', '110017', 'Origin',
        '2026-10-01', 'Drop', '400001', 'Destination', '2026-10-02'])
    buffer = io.BytesIO()
    workbook.save(buffer)
    for count in (1, 0):
        result = client.post('/admin/consignments/import', data={'file': (io.BytesIO(buffer.getvalue()), 'import.xlsx')}, follow_redirects=True)
        assert result.status_code == 200 and f'Added: {count}'.encode() in result.data
    shipment = client.get('/admin/consignments/list', query_string={'search': 'ER-EXCEL-001'}).get_json()['rows'][0]
    assert shipment['pickup_date'] == '2026-10-01' and shipment['drop_tag'] == 'Destination'
    with destination.begin() as connection:
        from uuid import UUID
        table = Base.metadata.tables['consignments']
        actual = connection.execute(select(table).where(table.c.id == UUID(shipment['id']))).mappings().one()
        assert actual['planned_pickup_date'].isoformat() == '2026-10-01'
