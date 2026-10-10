import csv
import io
import json
from datetime import date
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.admin.reporting import analyze, normalize_filters, parse_date
from app.models import Company, Consignment, MisReport, MisView, db


@pytest.fixture
def shipments(app):
    with app.app_context():
        client = Company(name='Acme & Sons')
        db.session.add(client)
        db.session.flush()
        db.session.add_all([
            Consignment(consignment_number='=SUM(1,2)', company=client, status='Delivered', pickup_date='2026-10-01', drop_date='2026-10-02', pieces=3, chargeable_weight='12.500', chargeable_volume='0.125', pickup_tag='Delhi', drop_tag='Mumbai'),
            Consignment(consignment_number='OPEN-1', company=client, status='In Transit', pickup_date='02/10/2026', drop_date='2026-10-03', pieces=2, chargeable_weight='7.250', pod_image='proof.png', invoice_file='invoice.pdf', pickup_tag='Delhi', drop_tag='Mumbai'),
            Consignment(consignment_number='UNDATED', status='mystery', pickup_date='bad date', pieces=1),
            Consignment(consignment_number='OLDER', status='Delivered', pickup_date='2026-09-30', pieces=4, chargeable_volume='1.500', pod_image='proof.png'),
        ])
        db.session.commit()
        return client.id


def test_shared_analytics_filters_dates_totals_and_estimated_flags(app, shipments):
    with app.app_context():
        filters = normalize_filters({'period': 'custom', 'start': '2026-10-01', 'end': '2026-10-09'}, today=date(2026,10,9))
        report = analyze(filters, today=date(2026,10,9))
        m = report['metrics']
        assert (m['total'], m['pieces'], m['weight'], m['volume']) == (2, 5, 19.75, .125)
        assert (m['delivered'], m['delivered_share'], m['open'], m['past_expected'], m['missing_pod'], m['missing_invoice'], m['excluded_undated']) == (1, 50, 1, 1, 1, 1, 1)
        assert (m['weight_recorded'],m['volume_recorded']) == (2,1)
        assert [item['count'] for item in report['trend']] == [1,1,0,0,0,0,0,0,0]
        assert report['clients'][0]['label'] == 'Acme & Sons'
        assert report['routes'] == [{'label':'Delhi → Mumbai','count':2}]
        all_time = analyze(normalize_filters({}))
        assert all_time['metrics']['total'] == 4 and all_time['metrics']['undated'] == 1
        assert analyze(normalize_filters({'company':'unassigned'}))['metrics']['total'] == 2
        assert analyze(normalize_filters({'company':str(shipments),'status':'Delivered'}))['metrics']['total'] == 1
        assert analyze(normalize_filters({'status':'Other'}))['metrics']['total'] == 1


@pytest.mark.parametrize('value,expected', [('2026-10-02',date(2026,10,2)),('02/10/2026',date(2026,10,2)),('02-10-2026',date(2026,10,2)),('2026-10-01T23:00:00Z',date(2026,10,2)),('2026-02-30',None),('2026-10-02garbage',None)])
def test_legacy_date_formats(value, expected):
    assert parse_date(value) == expected


def test_relative_filters_refresh_and_monthly_trend(app, shipments):
    with app.app_context():
        f = normalize_filters({'period':'month','start':'2026-09-01','end':'2026-09-30'}, today=date(2026,10,9))
        assert (f['start'],f['end']) == ('2026-10-01','2026-10-09')
        f = normalize_filters({'period':'quarter'},today=date(2026,5,11))
        assert (f['start'],f['end']) == ('2026-04-01','2026-05-11')
        report=analyze(normalize_filters({'period':'custom','start':'2026-01-01','end':'2026-12-31'}))
        assert report['grain']=='monthly' and len(report['trend'])==12
        assert sum(item['count'] for item in report['trend'])==3


@pytest.mark.parametrize('filters', [{'period':'unknown'}, {'period':'custom'}, {'period':'custom','start':'2026-10-09','end':'2026-10-01'}, {'company':'999'}, {'company':True}, {'status':['Delivered']}])
def test_invalid_filters_are_explicit_errors(admin_client, filters):
    response=admin_client.post('/admin/mis/reports',json={'name':'Invalid','filters':filters})
    assert response.status_code==400 and response.json['message']


@pytest.mark.parametrize('fmt', ['xlsx','csv','pdf'])
def test_exports_match_filtered_view_and_protect_user_text(admin_client, shipments, fmt):
    response=admin_client.get(f'/admin/mis/export?format={fmt}&company={shipments}&status=Delivered')
    assert response.status_code==200
    assert response.headers['Cache-Control']=='private, no-store'
    assert 'attachment;' in response.headers['Content-Disposition']
    if fmt=='xlsx':
        book=load_workbook(io.BytesIO(response.data))
        assert {'Management summary','Shipments','Status','Pickup trend','Clients','Routes','Attention'}<=set(book.sheetnames)
        rows=list(book['Shipments'].values)
        assert len(rows)==2 and rows[1][0]=='=SUM(1,2)'
        assert book['Shipments']['A2'].data_type=='s'
        assert len(book['Management summary']._charts)==3
    elif fmt=='csv':
        rows=list(csv.reader(io.StringIO(response.data.decode('utf-8-sig'))))
        assert len(rows)==2 and rows[1][0]=="'=SUM(1,2)"
    else:
        reader=PdfReader(io.BytesIO(response.data))
        text='\n'.join(page.extract_text() for page in reader.pages)
        assert 'Management MIS' in text and 'Acme & Sons' in text
        assert '12.500 kg' in text and 'no actual delivery time' in text
        assert any('/FontFile2' in font.get_object().get('/FontDescriptor',{}).get_object() for page in reader.pages for font in page['/Resources']['/Font'].values() if font.get_object().get('/FontDescriptor'))


def test_saved_views_crud_refresh_relative_dates_and_validate_columns(app, admin_client):
    data={'name':'Monthly review','notes':'Management','filters':{'period':'month'},'columns':['client','pieces'],'format':'xlsx'}
    response=admin_client.post('/admin/mis/views',json=data)
    assert response.status_code==200
    view_id=response.json['view']['id']
    assert admin_client.post('/admin/mis/views',json=data).status_code==400
    data.update(name='Monthly board review',format='pdf')
    assert admin_client.put(f'/admin/mis/views/{view_id}',json=data).json['view']['name']=='Monthly board review'
    assert b'Monthly board review' in admin_client.get('/admin/mis').data
    data['columns']=['unknown']
    assert admin_client.put(f'/admin/mis/views/{view_id}',json=data).status_code==400
    assert admin_client.delete(f'/admin/mis/views/{view_id}',json={}).status_code==200
    with app.app_context():
        assert MisView.query.count()==0


def test_report_snapshots_remain_immutable_manage_and_backup_every_file(app, admin_client, shipments):
    generated=admin_client.post('/admin/mis/reports',json={'name':'Board pack','notes':'October','format':'xlsx'}).json
    download=admin_client.get(generated['download_url'])
    before=download.data
    with app.app_context():
        row=db.session.get(MisReport,generated['id'])
        assert row.row_count==4 and json.loads(row.summary_json)['total']==4
        file_ref=row.file_ref
        Consignment.query.first().pieces=99
        db.session.commit()
    assert admin_client.get(generated['download_url']).data==before
    assert admin_client.put(f'/admin/mis/reports/{generated["id"]}',json={'name':'Updated display name','notes':'Retained snapshot'}).status_code==200
    assert b'Updated display name' in admin_client.get('/admin/mis').data
    # Missing PODs also stay explicitly reported; the saved MIS file must be linked in backup.
    backup=admin_client.get('/admin/generate-backup')
    archive=ZipFile(io.BytesIO(backup.data))
    data=json.loads(archive.read('data.json'))
    saved=data['mis_report'][0]
    assert saved['file_backup_status']=='included'
    assert saved['file_backup_path'].startswith('mis_reports/')
    assert archive.read(saved['file_backup_path'])==before
    assert admin_client.delete(f'/admin/mis/reports/{generated["id"]}',json={}).status_code==200
    assert not Path(app.instance_path,'uploads',file_ref).exists()
    assert admin_client.get(generated['download_url']).status_code==404
    with app.app_context():
        assert Consignment.query.count()==4


def test_report_failure_cleans_new_file_and_preserves_existing_history(app, admin_client, monkeypatch):
    def fail():
        raise RuntimeError('Database unavailable')
    with app.app_context():
        monkeypatch.setattr(db.session,'commit',fail)
        response=admin_client.post('/admin/mis/reports',json={'name':'Failure'})
        assert response.status_code==500
        assert MisReport.query.count()==0
        assert not list(Path(app.instance_path,'uploads').iterdir())


def test_mis_auth_and_cross_origin_mutations(client, admin_client):
    client.get('/admin/logout')
    assert client.get('/admin/mis').status_code==302
    assert client.post('/admin/mis/reports',json={'name':'Test'}).status_code==401
    client.post('/admin/login',data={'username':'admin','password':'test-admin-password'})
    assert client.post('/admin/mis/views',json={'name':'Test'},headers={'Origin':'https://other.example'}).status_code==403


def test_empty_reports_and_invalid_download_format(admin_client):
    assert admin_client.get('/admin/mis').status_code==200
    for fmt in ('xlsx','csv','pdf'):
        assert admin_client.get('/admin/mis/export?format='+fmt).status_code==200
    assert admin_client.get('/admin/mis/export?format=html').status_code==400
    assert admin_client.get('/admin/dashboard?period=custom').status_code==400


def test_invalid_measurements_and_large_date_spans_do_not_invent_totals(app):
    with app.app_context():
        db.session.add_all([
            Consignment(consignment_number='OLD', pickup_date='2000-01-01', chargeable_weight='-5', pieces=0),
            Consignment(consignment_number='NOW', pickup_date='2026-10-09', chargeable_weight='10', pieces=2),
        ])
        db.session.commit()
        report=analyze(normalize_filters({}))
        assert report['metrics']['weight']==10 and report['metrics']['pieces']==2
        assert report['metrics']['invalid_measurements']==1
        assert report['grain']=='yearly' and len(report['trend'])==27
        assert sum(item['count'] for item in report['trend'])==2


def test_report_snapshot_keeps_client_label_after_master_rename(app, admin_client, shipments):
    response=admin_client.post('/admin/mis/reports',json={'name':'Client pack','filters':{'company':str(shipments)}})
    assert response.status_code==200
    with app.app_context():
        db.session.get(Company,shipments).name='Renamed master'
        db.session.commit()
    html=admin_client.get('/admin/mis').get_data(as_text=True)
    assert 'Acme &amp; Sons · All statuses' in html


def test_existing_database_adds_mis_tables_without_changing_shipments(app, shipments):
    with app.app_context():
        MisReport.__table__.drop(db.engine)
        MisView.__table__.drop(db.engine)
        before=[(row.id,row.consignment_number,row.pieces) for row in Consignment.query.order_by(Consignment.id)]
    result=app.test_cli_runner().invoke(args=['upgrade-db'])
    assert result.exit_code==0 and 'mis_view table' in result.output and 'mis_report table' in result.output
    assert app.test_cli_runner().invoke(args=['check-db']).exit_code==0
    with app.app_context():
        assert [(row.id,row.consignment_number,row.pieces) for row in Consignment.query.order_by(Consignment.id)]==before
        assert MisReport.query.count()==0 and MisView.query.count()==0


def pdf_text(content):
    reader = PdfReader(io.BytesIO(content))
    assert reader.pages
    pages = [page.extract_text() or '' for page in reader.pages]
    assert all('GRAM SCS' in page for page in pages)
    assert 'Management MIS' in pages[0]
    return '\n'.join(pages)


def test_tracker_pdf_replaces_title_only_export_with_all_saved_shipments(admin_client, shipments):
    response = admin_client.get('/admin/consignments/export.pdf')
    assert response.status_code == 200
    text = pdf_text(response.data)
    assert 'Shipment register' in text
    for identifier in ('=SUM(1,2)', 'OPEN-1', 'UNDATED', 'OLDER'):
        assert identifier in text
    assert 'Chargeable weight (kg)' in text and '19.750 kg' in text


def test_all_pdf_entry_points_include_healthy_shipments_not_only_attention_records(app, admin_client):
    with app.app_context():
        db.session.add(Consignment(consignment_number='HEALTHY-PDF-001', status='Delivered', pieces=7, chargeable_weight=22.5, chargeable_volume=.75, pickup_date='2026-10-10', pickup_tag='Delhi', drop_tag='Mumbai', pod_image='pod.pdf', invoice_file='invoice.pdf'))
        db.session.commit()
    for url in ('/admin/mis/export?format=pdf', '/admin/consignments/export.pdf'):
        text = pdf_text(admin_client.get(url).data)
        assert 'HEALTHY-PDF-001' in text and 'No attention flags' in text
        assert '22.500' in text and '0.750' in text
    generated = admin_client.post('/admin/mis/reports', json={'name':'Saved PDF', 'format':'pdf'})
    assert generated.status_code == 200
    text = pdf_text(admin_client.get(generated.json['download_url']).data)
    assert 'HEALTHY-PDF-001' in text and 'Shipment register' in text


def test_pdf_empty_filters_explain_missing_results_instead_of_silent_zeroes(admin_client, shipments):
    response = admin_client.get('/admin/mis/export?format=pdf&period=custom&start=2027-01-01&end=2027-01-31')
    text = pdf_text(response.data)
    assert 'No shipments match these report filters.' in text
    assert 'database contains 4 saved shipments' in text
    assert 'pickup date is missing or invalid' in text
    assert '=SUM(1,2)' not in text and 'OPEN-1' not in text


def test_pdf_search_matches_tracker_and_preserves_selected_columns(app, admin_client, shipments):
    query = '?search=OPEN-1'
    table = admin_client.get('/admin/consignments/list'+query).json
    assert table['total'] == 1
    text = pdf_text(admin_client.get('/admin/consignments/export.pdf'+query).data)
    assert 'OPEN-1' in text and 'UNDATED' not in text and 'OLDER' not in text
    generated = admin_client.post('/admin/mis/reports', json={'name':'Custom columns','format':'pdf','filters':{'search':'OPEN-1'},'columns':['consignment_number','pickup_address','invoice']})
    assert generated.status_code == 200
    text = pdf_text(admin_client.get(generated.json['download_url']).data)
    assert 'Pickup address' in text and 'Invoice uploaded' in text and 'OPEN-1' in text
    with app.app_context():
        row = db.session.get(MisReport, generated.json['id'])
        assert json.loads(row.filters_json)['search'] == 'OPEN-1'
        assert row.row_count == 1


def test_pdf_register_paginates_and_keeps_last_row_with_long_escaped_addresses(app, admin_client):
    with app.app_context():
        for index in range(81):
            db.session.add(Consignment(consignment_number=f'PAGED-{index:03d}', status='Delivered', pickup_date='2026-10-10', pickup_address='A & B <Warehouse>\n' * (400 if index == 0 else 2), drop_address='Destination', pod_image='proof.pdf', invoice_file='invoice.pdf'))
        db.session.commit()
    response = admin_client.get('/admin/mis/export?format=pdf&columns=consignment_number&columns=pickup_address&columns=drop_address')
    assert response.status_code == 200
    reader = PdfReader(io.BytesIO(response.data))
    assert len(reader.pages) > 4
    text = pdf_text(response.data)
    assert 'PAGED-000' in text and 'PAGED-080' in text
    assert 'A & B <Warehouse>' in text


def test_empty_database_pdf_has_a_save_all_instruction(admin_client):
    text = pdf_text(admin_client.get('/admin/mis/export?format=pdf').data)
    assert 'No saved shipments are available to report.' in text
    assert 'Save All' in text


def test_corrupt_pdf_export_is_an_error_not_a_successful_download(admin_client, monkeypatch):
    from app.admin import mis_exports
    def corrupt(self, flowables, **kwargs):
        self.filename.write(b'not a PDF')
    monkeypatch.setattr(mis_exports.SimpleDocTemplate, 'build', corrupt)
    response = admin_client.get('/admin/mis/export?format=pdf')
    assert response.status_code == 500
    assert response.mimetype == 'application/json'
    assert 'could not be completed' in response.json['message']


def test_pdf_download_controls_are_available_everywhere(admin_client):
    assert b'id="shipment-pdf-export" data-report-download' in admin_client.get('/admin/consignments').data
    dashboard = admin_client.get('/admin/dashboard').get_data(as_text=True)
    assert 'data-report-download href="/admin/mis/export?' in dashboard
    assert 'report-downloads.js' in dashboard
